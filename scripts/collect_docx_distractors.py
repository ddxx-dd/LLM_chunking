"""docx_corpus 방해 문서 수집(2026-10-21): 질문이 없는 docx를 기존 QA 문서와 같은 출처·분포·기준으로 더 받는다.

  source env.sh
  .venv311/bin/python scripts/collect_docx_distractors.py plan       # 1. API 후보 목록(다운로드 없음)
  .venv311/bin/python scripts/collect_docx_distractors.py download   # 2. 후보 docx 받기(이미 받은 것은 건너뜀)
  .venv311/bin/python scripts/collect_docx_distractors.py select     # 3. 구조 기준·중복·정답 겹침 검사 → 칸마다 목표 수 고르기 → manifest에 덧붙임

출처: docxcorp.us(docx-corpus, ODC-By 1.0) API `https://api.docxcorp.us/documents`(신뢰도 높은 순 정렬, 쪽당 100개), 파일 `https://docxcorp.us/documents/<id>.docx`.
분포: 기존 QA 문서 94편의 (주제, 유형) 칸 비율대로 TARGET편을 나눈다(큰 나머지 방식). 칸마다 목표의 OVERSAMPLE배를 받는다.
기준(기존 94편을 같은 실측 함수 stats()로 잰 범위, 94편 모두 통과): 영어, API 신뢰도 ≥ 0.8, 단어 4,000~24,000, Word 앱(Microsoft Word),
  제목 스타일 문단 ≥ 6, 표 ≤ 3개·표 글자 비율 ≤ 0.18, 글상자·프레임 ≤ 2, 빈 문단 비율 ≤ 0.24, 문단 글자 수 중앙값 36~383, 그림 ≤ 22.
  각주는 기준에서 뺐다(기존 manifest의 n_footnotes = 0은 Docling이 docx 각주를 세지 못한 값이고, 실측하면 94편 중 36편에 각주가 있다).
중복: 기존 문서 id와 같으면 제외, 본문 8단어 묶음 Jaccard ≥ 0.3이면(기존·이미 고른 문서와) 제외.
정답 겹침: 기존 QA의 근거 인용(support)·정답이 방해 문서 본문에 나오면(8단어 묶음 절반 이상, 짧으면 정규화 문자열 포함) 그 문서를 뺀다.
저장: 원본 data/raw/docx_corpus/distractor/<주제>/<파일이름>_<id 앞 8자>.docx, 중간 기록 data/raw/docx_corpus/_meta/distractors/,
  manifest(data/manifests/docx_corpus.json)에 origin = "v3_distractor"로 덧붙인다(기존 항목은 건드리지 않음).
2026-10-21 실행 결과: 후보 3,899개(칸마다 목표의 20배) → 200편 통과(탈락 사유는 _meta/distractors/select_report.json).
  고르지 않은 후보 3,699개 docx는 지웠다(plan.json·select_report.json은 남김). 다시 받으려면 download를 다시 돌리면 된다.
"""
import argparse
import json
import random
import re
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from collections import Counter
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw" / "docx_corpus"
OUT = RAW / "distractor"
META = RAW / "_meta" / "distractors"
MANIFEST = ROOT / "data" / "manifests" / "docx_corpus.json"
QA = ROOT / "data" / "qa" / "docx_corpus.jsonl"
API = "https://api.docxcorp.us/documents"
FILE = "https://docxcorp.us/documents/{}.docx"
UA = {"User-Agent": "Mozilla/5.0 (research; DIP chunking lab)"}
TARGET, OVERSAMPLE, SEED = 200, 20, 0      # 기준 통과율이 6~9%(2026-10-21: 3배 → 53편, 12배 → 155편) → 20배로 늘림
WORDS = (4000, 24000)
MIN_CONF = 0.8


def get(url, binary=False, tries=4):
    for k in range(tries):
        try:
            r = urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=60).read()
            return r if binary else json.loads(r)
        except Exception as e:  # noqa: BLE001
            if k == tries - 1:
                raise
            print(f"  재시도 {k + 1}: {e}", flush=True)
            time.sleep(3 * (k + 1))


def quotas():
    """기존 QA 문서의 (주제, 유형) 칸 비율로 TARGET을 나눈다(큰 나머지)."""
    m = [e for e in json.loads(MANIFEST.read_text(encoding="utf-8")) if e.get("origin") != "v3_distractor"]
    c = Counter((e["topic"], e["type"]) for e in m)
    raw = {k: v * TARGET / len(m) for k, v in c.items()}
    q = {k: int(v) for k, v in raw.items()}
    for k in sorted(raw, key=lambda k: raw[k] - q[k], reverse=True)[:TARGET - sum(q.values())]:
        q[k] += 1
    return q, {e["id"] for e in m}


def plan():
    q, have = quotas()
    rnd = random.Random(SEED)
    old = META / "plan.json"
    cands = json.loads(old.read_text(encoding="utf-8"))["candidates"] if old.exists() else {}     # 이미 받은 후보는 유지하고 덧붙임
    for (topic, typ), n in sorted(q.items()):
        need = n * OVERSAMPLE - sum(x["cell"] == [topic, typ] for x in cands.values())
        if need <= 0:
            continue
        p = dict(lang="en", topic=topic, type=typ, min_confidence=MIN_CONF, limit=100, page=1)
        first = get(API + "?" + urllib.parse.urlencode(p))
        pages = list(range(1, first["pages"] + 1))
        rnd.shuffle(pages)
        got = []
        for pg in pages:
            d = first if pg == 1 else get(API + "?" + urllib.parse.urlencode({**p, "page": pg}))
            ok = [x for x in d["documents"] if x["language"] == "en" and WORDS[0] <= x["word_count"] <= WORDS[1]
                  and x["id"] not in have and x["id"] not in cands]
            rnd.shuffle(ok)
            got += ok
            time.sleep(0.3)
            if len(got) >= need:
                break
        for x in got[:need]:
            cands[x["id"]] = {**x, "cell": [topic, typ]}
        print(f"[{topic} / {typ}] 목표 {n}, 후보 {min(len(got), need)} (전체 {first['total']}건, {first['pages']}쪽)", flush=True)
    META.mkdir(parents=True, exist_ok=True)
    (META / "plan.json").write_text(json.dumps({"quotas": {f"{a}|{b}": n for (a, b), n in q.items()}, "candidates": cands},
                                               ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"후보 {len(cands)}개 → {META / 'plan.json'}")


def stem_of(x):
    s = re.sub(r"\.docx$", "", x["filename"], flags=re.I)
    s = re.sub(r"[^A-Za-z0-9._-]+", "-", s).strip("-.")[:60] or "doc"
    return f"{s}_{x['id'][:8]}"


def download():
    P = json.loads((META / "plan.json").read_text(encoding="utf-8"))
    fails = {}
    for i, (did, x) in enumerate(P["candidates"].items()):
        path = OUT / x["cell"][0] / f"{stem_of(x)}.docx"
        if path.exists():
            continue
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            b = get(FILE.format(did), binary=True)
            if not b.startswith(b"PK"):
                raise ValueError("docx(zip)가 아님")
            path.write_bytes(b)
        except Exception as e:  # noqa: BLE001
            fails[did] = str(e)
        time.sleep(0.3)
        if i % 50 == 0:
            print(f"  {i}/{len(P['candidates'])}", flush=True)
    (META / "download_fail.json").write_text(json.dumps(fails, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"다운로드 실패 {len(fails)}개")


def stats(path):
    """python-docx로 구조 실측(기존 manifest 항목과 같은 이름)."""
    import docx
    d = docx.Document(str(path))
    paras = d.paragraphs
    texts = [p.text for p in paras]
    head = sum(1 for p in paras if p.style is not None and re.match(r"(Heading|Title)\b", p.style.name or "") and p.text.strip())
    fake = sum(1 for p in paras if p.text.strip() and len(p.text) < 120 and p.runs and all(r.bold for r in p.runs if r.text.strip())
               and not (p.style is not None and re.match(r"(Heading|Title)\b", p.style.name or "")))
    tchars = sum(len(c.text) for t in d.tables for row in t.rows for c in row.cells)
    body = sum(len(t) for t in texts)
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        doc_xml = z.read("word/document.xml").decode("utf-8", "ignore")
        app = z.read("docProps/app.xml").decode("utf-8", "ignore") if "docProps/app.xml" in names else ""
        foot = z.read("word/footnotes.xml").decode("utf-8", "ignore") if "word/footnotes.xml" in names else ""
    tag = lambda t: (re.search(rf"<{t}>([^<]*)</{t}>", app) or [None, ""])[1]
    nonempty = [len(t) for t in texts if t.strip()]
    return {"application": tag("Application"), "pages_app": int(tag("Pages") or 0), "heading_count": head,
            "table_count": len(d.tables), "table_char_ratio": round(tchars / max(1, tchars + body), 4),
            "fake_bold_paragraphs": fake, "empty_para_ratio": round(sum(1 for t in texts if not t.strip()) / max(1, len(texts)), 4),
            "median_para_chars": int(np.median(nonempty)) if nonempty else 0,
            "textbox_count": doc_xml.count("<w:txbxContent"), "frame_count": doc_xml.count("<w:framePr"),
            "image_count": sum(n.startswith("word/media/") for n in names),
            "footnote_count": max(0, len(re.findall(r"<w:footnote [^>]*w:id=", foot)) - 2) if foot else 0,
            "docx_total_chars": tchars + body}, "\n".join(texts + [c.text for t in d.tables for row in t.rows for c in row.cells])


def fails_rules(s):
    r = []
    if "Microsoft" not in s["application"] or "Word" not in s["application"]:
        r.append("Word 앱 아님")
    if s["heading_count"] < 6:
        r.append("제목 < 6")
    if s["table_count"] > 3 or s["table_char_ratio"] > 0.18:
        r.append("표 많음")
    if s["textbox_count"] > 2 or s["frame_count"] > 2:
        r.append("글상자·프레임")
    if s["empty_para_ratio"] > 0.24:
        r.append("빈 문단 많음")
    if not 36 <= s["median_para_chars"] <= 383:
        r.append("문단 길이")
    if s["image_count"] > 22:
        r.append("그림 많음")
    return r


def norm(t):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]+", " ", t.lower())).strip()


def grams(t, n=8):
    w = norm(t).split()
    return {" ".join(w[i:i + n]) for i in range(len(w) - n + 1)}


def select():
    import docx  # noqa: F401
    P = json.loads((META / "plan.json").read_text(encoding="utf-8"))
    q = {tuple(k.split("|")): n for k, n in P["quotas"].items()}
    man = json.loads(MANIFEST.read_text(encoding="utf-8"))
    old = [e for e in man if e.get("origin") != "v3_distractor"]
    if len(old) != len(man):
        sys.exit("manifest에 이미 v3_distractor가 있다: 중복 추가 방지로 멈춤")
    # 기존 문서 본문 묶음(중복 검사)과 QA 근거(정답 겹침 검사)
    print("기존 문서 읽는 중", flush=True)
    seen = []
    for e in old:
        _, t = stats(ROOT / e["path"])
        seen.append(grams(t))
    qa = [json.loads(l) for l in QA.read_text(encoding="utf-8").splitlines() if l.strip()]
    probes = []
    for x in qa:
        for s in (x["evidence"].get("support") or "", x["answer"]):
            g = grams(s)
            if g:
                probes.append((x["qid"], g, None))
            elif len(norm(s)) >= 30:
                probes.append((x["qid"], None, norm(s)))
    rnd = random.Random(SEED)
    by_cell = {}
    for did, x in P["candidates"].items():
        by_cell.setdefault(tuple(x["cell"]), []).append((did, x))
    report, chosen, rej = [], [], Counter()
    for cell in sorted(q):
        items = by_cell.get(cell, [])
        rnd.shuffle(items)
        n_ok = 0
        for did, x in items:
            if n_ok >= q[cell]:
                break
            path = OUT / cell[0] / f"{stem_of(x)}.docx"
            if not path.exists():
                rej["다운로드 실패"] += 1; continue
            try:
                s, text = stats(path)
            except Exception as e:  # noqa: BLE001
                rej["열기 실패"] += 1; report.append((did, f"열기 실패: {e}")); continue
            bad = fails_rules(s)
            if bad:
                rej.update(bad); report.append((did, ", ".join(bad))); continue
            g = grams(text)
            dup = max((len(g & h) / max(1, len(g | h)) for h in seen), default=0)
            if dup >= 0.3:
                rej["중복(본문 겹침)"] += 1; report.append((did, f"중복 {dup:.2f}")); continue
            nt = norm(text)
            leak = sorted({qid for qid, pg, ps in probes if (pg is not None and len(pg & g) >= 0.5 * len(pg)) or (ps is not None and ps in nt)})
            if leak:
                rej["정답 겹침"] += 1; report.append((did, "정답 겹침: " + ", ".join(leak[:5]))); continue
            seen.append(g)
            n_ok += 1
            chosen.append({"file": path.name, "path": str(path.relative_to(ROOT)), "lang": "en", "type": cell[1], "topic": cell[0],
                           "id": did, "orig_filename": x["filename"], "url": FILE.format(did), "word_count_meta": x["word_count"],
                           "confidence": round(x["confidence"], 4), **s, "origin": "v3_distractor"})
        print(f"[{cell[0]} / {cell[1]}] 목표 {q[cell]}, 고름 {n_ok}, 후보 {len(items)}", flush=True)
    short = TARGET - len(chosen)
    if short > 0:
        print(f"목표보다 {short}편 부족: 모자란 주제의 다른 유형 → 그 밖의 남은 후보 순서로 채움", flush=True)
        used = {c["id"] for c in chosen}
        lack = Counter()
        for cell, n in q.items():
            lack[cell[0]] += n - sum(c["topic"] == cell[0] and c["type"] == cell[1] for c in chosen)
        rest = [(d, x) for d, x in P["candidates"].items() if d not in used]
        rnd.shuffle(rest)
        rest.sort(key=lambda dx: -lack[dx[1]["cell"][0]])
        for did, x in rest:
            if short == 0:
                break
            path = OUT / x["cell"][0] / f"{stem_of(x)}.docx"
            if not path.exists():
                continue
            try:
                s, text = stats(path)
            except Exception:  # noqa: BLE001
                continue
            if fails_rules(s):
                continue
            g = grams(text)
            if max((len(g & h) / max(1, len(g | h)) for h in seen), default=0) >= 0.3:
                continue
            nt = norm(text)
            if any((pg is not None and len(pg & g) >= 0.5 * len(pg)) or (ps is not None and ps in nt) for _, pg, ps in probes):
                continue
            seen.append(g); short -= 1; lack[x["cell"][0]] -= 1
            chosen.append({"file": path.name, "path": str(path.relative_to(ROOT)), "lang": "en", "type": x["cell"][1], "topic": x["cell"][0],
                           "id": did, "orig_filename": x["filename"], "url": FILE.format(did), "word_count_meta": x["word_count"],
                           "confidence": round(x["confidence"], 4), **s, "origin": "v3_distractor"})
    ids = {Path(e["file"]).stem for e in old}
    assert not ids & {Path(c["file"]).stem for c in chosen}, "doc_id 충돌"
    (META / "select_report.json").write_text(json.dumps({"rejected": dict(rej), "detail": report}, ensure_ascii=False, indent=1), encoding="utf-8")
    MANIFEST.write_text(json.dumps(old + chosen, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"고름 {len(chosen)}편 → manifest {len(old)} + {len(chosen)} = {len(old) + len(chosen)}편 | 탈락 사유 {dict(rej)}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("step", choices=["plan", "download", "select"])
    {"plan": plan, "download": download, "select": select}[ap.parse_args().step]()
