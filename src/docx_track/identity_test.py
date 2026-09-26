"""항등 테스트(docx, smart_chunk.md 8절/10절 2번) - elements의 "자기 자신 text"를 자기
"loc"에 그대로 되돌려 써서, 원본과 비교한다. 7절 병합 writer의 문단 해석 로직(paraId/
paraIds/table_index fallback)이 실제로 정확한 위치를 찾는지 검증하는 게 목적 - 번역이
아니라 "원문 되돌려쓰기"라서 텍스트 자체는 파싱 때 이미 검증됨(parse_open_issues.md 체크1).

★ 실측 확인(2026-09-26) - 처음엔 run.text 대입(+ iter_inner_content) 방식으로 썼는데,
그룹 도형(wgp) 여러 개를 포함한 문서 전체를 되돌려 쓰면 그룹 도형 안 문단 34개가
통째로 사라지는 문제가 있었다. 원인: 그룹 도형을 호스팅하는 "앵커 문단"(예: "변화하는
시대, 변화하는 구매자")의 run 중 하나는 텍스트가 아니라 w:drawing/mc:AlternateContent만
담고 있는데, 그 run에 `.text = "..."`를 대입하면 python-docx가 run 내용 전체를
지우고(clear_content) 새 w:t만 넣어서 도형 자체가 삭제된다(통제 실험으로 확인: 이
앵커 문단 "하나만" 되돌려 써도 재현됨). 그래서 run 단위가 아니라 **문단에 직접 속한
w:t만** 골라 고치는 방식으로 바꿨다(_direct_text_nodes) - 조건: t.iterancestors(w:p)의
첫 번째가 이 문단 자신일 때만(중첩된 도형/텍스트박스 안 w:t는 그 안쪽 문단이 따로
처리하므로 제외)."""
import sys
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.path.append(str(Path(__file__).resolve().parent.parent))
from docx_track.parse import add_para_ids, parse_docx
from docx_track.writer import build_paraid_index, write_element_text


def _write_own_text_back(anchored_path, elements, out_path):
    """7절 docx writer 로직(paraId 우선, paraIds는 첫 문단에 쓰고 나머진 비움, table_index
    fallback)을 그대로 따르되 번역문 대신 요소 자기 자신의 text를 쓴다. 실제 쓰기는
    docx_track/writer.py의 w:t 직접 수정 방식(위 모듈 docstring)을 그대로 씀 -
    merge_checks.py의 표시 테스트와 로직을 공유."""
    doc = DocxDocument(str(anchored_path))
    paraid_to_p = build_paraid_index(doc)

    counts = {"written": 0, "skipped_no_loc": 0, "skipped_flag": 0, "skipped_no_text_node": 0}

    def record(result):
        if result == "ok":
            counts["written"] += 1
        elif result == "no_text_node":
            counts["skipped_no_text_node"] += 1
        else:
            counts["skipped_no_loc"] += 1

    for e in elements:
        if e["label"] == "table":
            for c in e["cells"]:
                if c.get("skip") or e.get("skip"):  # 실제 병합에서도 안 건드릴 대상(5절) - 원문 그대로 둠
                    counts["skipped_flag"] += 1
                    continue
                record(write_element_text(doc, paraid_to_p, c.get("loc"), c["text"]))
            continue
        if e["label"] == "picture" or e.get("skip"):
            counts["skipped_flag"] += 1
            continue
        record(write_element_text(doc, paraid_to_p, e.get("loc"), e["text"], e.get("auto_num")))

    doc.save(str(out_path))
    return counts


def _in_fallback(p_xml):
    """★ 실측 확인(2026-09-26, 항등 테스트에서 발견): 도형/텍스트박스가 있는 문단은
    mc:AlternateContent 안에 mc:Choice(최신 DrawingML)와 mc:Fallback(구버전 VML) 두
    가지로 "같은 내용"이 중복 저장돼 있을 수 있다 - 실제로 렌더링되는 건 Choice 쪽 하나뿐인데
    body.iter(w:p)로 순수하게 순회하면 Fallback 쪽 문단까지 별개 문단처럼 두 번 잡힌다.
    Docling은 Fallback을 무시하고 Choice만 읽으므로, 검증 시에도 Fallback 문단은
    제외해야 원본과 공정하게 비교된다."""
    return any(a.tag.rsplit("}", 1)[-1] == "Fallback" for a in p_xml.iterancestors())


def _full_text(docx_path):
    """문서 전체 텍스트(본문 + 표 셀 포함, body.iter(w:p)가 둘 다 잡음)를 공백 무시하고
    이어붙인다. mc:Fallback 안 중복 문단은 제외(위 _in_fallback 참고)."""
    doc = DocxDocument(str(docx_path))
    parts = [Paragraph(p, doc).text for p in doc.element.body.iter(qn("w:p")) if not _in_fallback(p)]
    return "".join("".join(t.split()) for t in parts)


def _count_bold_italic(docx_path):
    doc = DocxDocument(str(docx_path))
    bold = italic = 0
    for r_xml in doc.element.body.iter(qn("w:r")):
        rpr = r_xml.find(qn("w:rPr"))
        if rpr is None:
            continue
        if rpr.find(qn("w:b")) is not None:
            bold += 1
        if rpr.find(qn("w:i")) is not None:
            italic += 1
    return bold, italic


def _count_structural(docx_path):
    """항등 테스트 구조 보존 기준(3번 요청) - 저장 전후 w:p/w:drawing/w:pict 개수가
    같아야 통과. 그룹 도형 소실 버그를 다시 놓치지 않기 위한 회귀 감지용."""
    doc = DocxDocument(str(docx_path))
    body = doc.element.body
    return {
        "w:p": sum(1 for _ in body.iter(qn("w:p"))),
        "w:drawing": sum(1 for _ in body.iter(qn("w:drawing"))),
        "w:pict": sum(1 for _ in body.iter(qn("w:pict"))),
    }


def run_identity_test(src_path):
    src_path = Path(src_path)
    anchored_path = src_path.with_suffix(".anchored.docx")
    if not anchored_path.exists():
        anchored_path = add_para_ids(src_path)
    elements = parse_docx(src_path)

    out_path = src_path.with_name(src_path.stem + ".identitytest.docx")
    counts = _write_own_text_back(anchored_path, elements, out_path)

    orig_text = _full_text(src_path)
    new_text = _full_text(out_path)
    text_match = orig_text == new_text

    b0, i0 = _count_bold_italic(src_path)
    b1, i1 = _count_bold_italic(out_path)

    struct_before = _count_structural(anchored_path)
    struct_after = _count_structural(out_path)
    struct_match = struct_before == struct_after

    print(f"[identity docx] {src_path.name}")
    print(f"  쓴 요소/셀 수: {counts['written']}, loc 없어서 건너뜀: {counts['skipped_no_loc']}, "
          f"skip 표시라 원문 유지: {counts['skipped_flag']}, "
          f"문단에 직접 속한 w:t가 없어서 건너뜀(no_text_node): {counts['skipped_no_text_node']}")
    print(f"  텍스트(공백 무시) 완전 일치: {text_match}")
    if not text_match:
        print(f"    원본 길이 {len(orig_text)}자, 재작성 길이 {len(new_text)}자")
        n = min(len(orig_text), len(new_text))
        for i in range(n):
            if orig_text[i] != new_text[i]:
                print(f"    첫 불일치 위치 {i}: 원본 ...{orig_text[max(0, i - 20):i + 20]}...")
                print(f"                    재작성 ...{new_text[max(0, i - 20):i + 20]}...")
                break
    print(f"  구조 보존(w:p/w:drawing/w:pict 개수 저장 전후 동일): {struct_match}"
          f" (전: {struct_before}, 후: {struct_after})")
    print(f"  bold run 개수: 원본 {b0} -> 재작성 {b1} (차이 {b1 - b0})")
    print(f"  italic run 개수: 원본 {i0} -> 재작성 {i1} (차이 {i1 - i0})")

    passed = text_match and struct_match
    print(f"  ▶ 항등 테스트 {'통과' if passed else '실패'}")
    return {"text_match": text_match, "struct_match": struct_match, "passed": passed,
            **counts, "bold": (b0, b1), "italic": (i0, i1)}


if __name__ == "__main__":
    for f in sys.argv[1:]:
        run_identity_test(f)
