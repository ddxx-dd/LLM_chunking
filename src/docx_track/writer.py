"""docx 병합 writer 핵심(smart_chunk.md 7절) - "문단에 직접 속한 w:t만 고친다" 방식
(그룹 도형을 run.text 대입으로 파괴하던 버그의 해결책, identity_test.py에서 실측 확인).
identity_test.py(자기 텍스트 되돌려쓰기)와 merge_checks.py(마킹 테스트)가 이 모듈을
공유한다."""
from docx.oxml.ns import qn


def _local(tag):
    return tag.rsplit("}", 1)[-1]


def direct_text_nodes(para_xml):
    """para_xml에 직접 속한 w:t만 반환 - 조건: t.iterancestors(w:p)의 첫 번째가
    이 문단 자신일 때만(중첩된 도형/텍스트박스 안 w:t는 제외)."""
    return [t for t in para_xml.iter(qn("w:t"))
            if next(t.iterancestors(qn("w:p")), None) is para_xml]


def build_paraid_index(doc):
    return {p.get(qn("w14:paraId")): p for p in doc.element.body.iter(qn("w:p"))}


def resolve_targets(doc, paraid_to_p, loc):
    if "paraId" in loc:
        p = paraid_to_p.get(loc["paraId"])
        return [p] if p is not None else []
    if "paraIds" in loc:
        return [paraid_to_p[pid] for pid in loc["paraIds"] if pid in paraid_to_p]
    cell = doc.tables[loc["table_index"]].cell(loc["row"], loc["col"])
    return [p._p for p in cell.paragraphs]


def fallback_sibling(para_xml):
    """★ 실측 확인(2026-09-26, 표시 테스트에서 발견) - mc:AlternateContent의 Choice(최신
    DrawingML) 쪽 문단만 Docling 요소가 되고 Fallback(구버전 VML) 쪽은 요소가 없어서
    한 번도 안 써진다(원문 그대로 남음 - 번역 후엔 원문·번역문이 섞인 상태가 됨). Choice
    쪽 문단이면 같은 AlternateContent 안 Fallback 쪽의 "같은 순번"(각 branch 안에서
    문서 순서 인덱스) 문단을 찾아 반환한다(없으면 None) - 실측: 인용구 카드 문단들이
    Choice/Fallback 양쪽에 순서대로 1:1 대응하는 것으로 확인됨."""
    choice = next((a for a in para_xml.iterancestors() if _local(a.tag) == "Choice"), None)
    if choice is None:
        return None
    alt = choice.getparent()
    if alt is None or _local(alt.tag) != "AlternateContent":
        return None
    fallback = next((c for c in alt if _local(c.tag) == "Fallback"), None)
    if fallback is None:
        return None
    choice_ps = list(choice.iter(qn("w:p")))
    try:
        idx = choice_ps.index(para_xml)
    except ValueError:
        return None
    fallback_ps = list(fallback.iter(qn("w:p")))
    if idx >= len(fallback_ps):
        return None
    return fallback_ps[idx]


def write_paragraph_text(para_xml, text, mirror_fallback=True):
    """이 문단(과, mirror_fallback이면 대응하는 Fallback 문단에도) 직접 속한 w:t에 text를
    쓴다. 반환: "ok" | "no_text_node"."""
    targets = [para_xml]
    if mirror_fallback:
        sib = fallback_sibling(para_xml)
        if sib is not None:
            targets.append(sib)

    wrote_any = False
    for p in targets:
        t_nodes = direct_text_nodes(p)
        if not t_nodes:
            continue
        t_nodes[0].text = text
        t_nodes[0].set(qn("xml:space"), "preserve")
        for extra in t_nodes[1:]:
            extra.text = ""
        wrote_any = True
    return "ok" if wrote_any else "no_text_node"


def write_element_text(doc, paraid_to_p, loc, text, auto_num=None, mirror_fallback=True):
    """요소/셀 하나의 loc에 text를 쓴다(paraId 우선, paraIds는 첫 문단만 쓰고 나머진 비움,
    table_index fallback). auto_num이 있으면 번역문 앞에서 뗀다(3절)."""
    if not loc:
        return "no_loc"
    if auto_num and text.startswith(auto_num):
        text = text[len(auto_num):]
    targets = resolve_targets(doc, paraid_to_p, loc)
    if not targets:
        return "no_loc"
    result = "ok"
    for i, para_xml in enumerate(targets):
        text_for_this_para = text if i == 0 else ""
        r = write_paragraph_text(para_xml, text_for_this_para, mirror_fallback=mirror_fallback)
        if i == 0 and r == "no_text_node":
            result = "no_text_node"
    return result
