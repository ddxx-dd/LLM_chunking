"""docx -> elements 파서 (Docling MsWordDocumentBackend, paraId 앵커링). smart_chunk.md 3절.

실측 확인(2026-09-26, 커머스 문서로 검증): backend.paragraph_to_items는 Docling 자체
댓글-연결 기능의 부산물이라 "서식이 섞이지 않은 단일 run 문단"만 커버한다 - 소스 확인 결과
_handle_text_elements()의 list_item 분기가 항상 빈 elem_ref로 일찍 return해서(2216행)
paragraph_to_items 등록 자체를 못 하고, 서식 혼합 문단도 마찬가지로 안 잡힌다. 그래서
AnchoredWordBackend로 _handle_text_elements 자체를 감싸서(래핑) 호출 전후 doc.texts
길이를 비교, 그 호출에서 새로 생긴 텍스트 객체마다 obj_to_paraid[id(객체)] = 그 문단의
paraId를 직접 기록한다 - list_item이든 서식 혼합이든 내부적으로 해당 메서드 한 번의
호출 안에서 doc.texts에 append되므로 100% 잡힌다(paragraph_to_items가 내부적으로
뭘 등록하고 안 하는지와 무관). self_ref가 아니라 객체 id()를 키로 쓰는 이유: convert()
뒷단에서 furniture 삭제 등으로 self_ref 문자열이 재계산될 수 있어(리스트 재정렬), 문자열
self_ref는 변할 수 있지만 같은 파이썬 객체(id)는 그대로 유지된다.
표 셀은 이 문제와 무관 - python-docx 문단 XML에서 w14:paraId를 직접 읽으므로 전부 커버됨
(실측: 병합 셀 포함 8/8 셀 모두 paraId 확인)."""
import random
import sys
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from docling.backend.msword_backend import MsWordDocumentBackend
from docling.datamodel.base_models import InputFormat
from docling.datamodel.document import InputDocument

sys.path.append(str(Path(__file__).resolve().parent.parent))
from docobj import NUM_ONLY, build_table_markdown

FURNITURE_LABELS = {"page_header", "page_footer"}


class AnchoredWordBackend(MsWordDocumentBackend):
    """_handle_text_elements를 감싸서 obj_to_paraid[id(text_item)] = paraId를 직접 기록한다
    (위 모듈 docstring 참고) - paragraph_to_items의 커버리지 구멍(list_item, 서식 혼합
    문단)을 우회하는 실측 확인된 방법."""

    def __init__(self, in_doc, path_or_stream, options=None):
        super().__init__(in_doc, path_or_stream, options=options)
        self.obj_to_paraid = {}

    def _handle_text_elements(self, element, doc, skip_empty_text=False):
        before = len(doc.texts)
        result = super()._handle_text_elements(element, doc, skip_empty_text=skip_empty_text)
        para_id = element.get(qn("w14:paraId"))
        if para_id:
            for t in doc.texts[before:]:
                self.obj_to_paraid[id(t)] = para_id
        return result


def add_para_ids(src_path):
    """모든 <w:p>에 w14:paraId가 없으면 부여하고 *.anchored.docx로 저장(3절 순서 1번).
    python-docx의 nsmap에 w14가 이미 등록돼 있어 별도 네임스페이스 등록 불필요(실측 확인).
    ★ 실측 확인(2026-09-26, 항등 테스트에서 발견) - 이미 anchored 파일이 있으면 그대로
    재사용한다(멱등성 필수) - 매번 새로 만들면 실행할 때마다 다른 무작위 paraId가 부여돼서,
    캐싱된 elements JSON의 paraId가 최신 anchored 파일과 어긋나는 버그가 실제로 발생했다
    (parse → identity_test처럼 add_para_ids가 두 번 이상 호출되는 경로에서 특히 치명적 -
    §0 원칙 4번 "문서 하나당 파싱은 1번만, JSON 캐싱"과도 직결됨)."""
    src_path = Path(src_path)
    anchored_path = src_path.with_suffix(".anchored.docx")
    if anchored_path.exists():
        return anchored_path
    doc = DocxDocument(str(src_path))
    seen = {p.get(qn("w14:paraId")) for p in doc.element.body.iter(qn("w:p")) if p.get(qn("w14:paraId"))}
    for p in doc.element.body.iter(qn("w:p")):
        if p.get(qn("w14:paraId")):
            continue
        new_id = "%08X" % random.randint(1, 0x7FFFFFFF)
        while new_id in seen:
            new_id = "%08X" % random.randint(1, 0x7FFFFFFF)
        p.set(qn("w14:paraId"), new_id)
        seen.add(new_id)
    doc.save(str(anchored_path))
    return anchored_path


def _has_descendant(xml_elem, tag):
    return next(xml_elem.iter(tag), None) is not None


def _skip_reason_for_paragraph(xml_elem):
    """skip 규칙(3절) - furniture(page_header/footer)는 이미 순회 단계에서 제외되므로 여기 없음."""
    if xml_elem is None:
        return None
    if _has_descendant(xml_elem, qn("w:fldChar")) or _has_descendant(xml_elem, qn("w:instrText")):
        return "field_code"
    if _has_descendant(xml_elem, qn("w:footnoteReference")):
        return "footnote_ref"
    if _has_descendant(xml_elem, qn("w:ins")) or _has_descendant(xml_elem, qn("w:del")):
        return "tracked_change"
    return None


def _detect_auto_num(docling_text, raw_xml_elem, docx_obj):
    """Docling이 Word 자동번호매김 제목에 번호를 합성해서 붙이는 경우(_add_heading()의
    numbered_headers 카운터, 실측 확인 - "1 미래 전략" 등) 원본 문단의 리터럴 텍스트에는
    이 번호가 없다. Docling 텍스트가 원본 문단 텍스트를 접미사로 가지면 그 차이를
    auto_num으로 뗀다 - 병합 시 번역문 앞에 이 번호를 그대로 쓰면 안 되므로(3절/7절)."""
    if raw_xml_elem is None:
        return None
    raw_text = Paragraph(raw_xml_elem, docx_obj).text
    if raw_text and docling_text != raw_text and docling_text.endswith(raw_text):
        return docling_text[:len(docling_text) - len(raw_text)]
    return None


def _build_table_element(item, eid, docx_tables, table_index):
    """표 = 요소 1개 + cells 목록(6절 두 층 구조). 셀 loc은 python-docx 표에서 직접 읽은
    paraId가 1순위 - Docling 표 순서와 python-docx doc.tables 순서가 같다는 가정을
    (0,0) 셀 텍스트 일치로 검증하고, 다르면 table_index/row/col로 폴백(3절)."""
    n_rows, n_cols = item.data.num_rows, item.data.num_cols
    raw_cells = list(item.data.table_cells)
    cell_texts = {}
    for c in raw_cells:
        r, col = c.start_row_offset_idx, c.start_col_offset_idx
        cell_texts[(r, col)] = (c.text or "").strip()
    md, span_map = build_table_markdown(n_rows, n_cols, cell_texts)

    pt = docx_tables[table_index] if table_index < len(docx_tables) else None
    mismatch = pt is None
    if not mismatch:
        try:
            mismatch = pt.cell(0, 0).text.strip() != cell_texts.get((0, 0), "")
        except IndexError:
            mismatch = True

    cells = []
    for cid, c in enumerate(raw_cells):
        r, col = c.start_row_offset_idx, c.start_col_offset_idx
        text = cell_texts[(r, col)]
        if mismatch:
            loc = {"table_index": table_index, "row": r, "col": col}
        else:
            para_ids = [p._p.get(qn("w14:paraId")) for p in pt.cell(r, col).paragraphs]
            loc = {"paraIds": para_ids} if para_ids else {"table_index": table_index, "row": r, "col": col}
        cell = {
            "cell_id": cid, "row": r, "col": col,
            "row_span": c.row_span, "col_span": c.col_span,
            "header": bool(c.column_header or c.row_header),
            "text": text, "span_in_table": span_map[(r, col)], "loc": loc,
        }
        if NUM_ONLY.fullmatch(text):
            cell["skip"] = "num_only"
        cells.append(cell)

    return {
        "id": eid, "label": "table", "level": None, "text": md, "loc": None,
        "skip": "table_mismatch" if mismatch else None,
        "table_id": table_index, "caption_ids": [], "n_rows": n_rows, "n_cols": n_cols,
        "cells": cells,
    }


def parse_docx(src_path):
    anchored_path = add_para_ids(src_path)
    in_doc = InputDocument(path_or_stream=anchored_path, format=InputFormat.DOCX, backend=AnchoredWordBackend)
    backend = in_doc._backend
    doc = backend.convert()

    # field_code/footnote_ref/tracked_change skip 판정용 - paragraph_to_items가 아니라
    # anchored 문서의 XML을 직접 순회해서 만든다(위 obj_to_paraid와 같은 이유: paraId
    # 위치는 100% 잡아도 이 dict를 paragraph_to_items에서 만들면 여전히 커버리지 구멍이
    # 생김 - 실측: TOC 페이지번호 필드코드 문단이 paragraph_to_items엔 없어서 skip 판정을
    # 놓치는 걸 항등 테스트에서 발견함).
    paraid_to_xmlelem = {p.get(qn("w14:paraId")): p for p in backend.docx_obj.element.body.iter(qn("w:p"))}

    # ★ 실측 확인(2026-09-26, 항등 테스트에서 발견) - doc.iterate_items()가 "표 셀 안 문단은
    # 별도 top-level 아이템으로 안 내놓는다"는 가정(3절 순서 5번)이 틀렸다: 표 하나에서
    # 6개 셀 문단이 표 요소와 별개로 "text" 아이템으로도 다시 나와서, 같은 paraId를 가진
    # 요소가 2개(표 셀 + 일반 text) 생기고 병합 시 나중 것이 먼저 것을 덮어씀. 그래서 표
    # 셀 문단의 paraId 집합을 미리 만들어두고, 일반 텍스트 루프에서 명시적으로 제외한다.
    table_paraids = {pid for pid, p in paraid_to_xmlelem.items()
                     if any(a.tag == qn("w:tbl") for a in p.iterancestors())}

    docx_tables = backend.docx_obj.tables  # Docling이 내부적으로 연 python-docx 문서 재사용(같은 파일 두 번 안 엶)

    elements = []
    eid = 0
    table_index = 0
    pending = None  # 인라인 그룹(서식 섞인 한 문단이 여러 조각으로 쪼개진 경우) 누적용(3절 순서 4번)

    def flush():
        nonlocal pending, eid
        if pending is None:
            return
        text = "".join(pending["text_parts"]).strip()
        if text:
            para_id = pending["para_id"]
            raw_xml_elem = paraid_to_xmlelem.get(para_id)
            skip = "formula" if pending["label"] == "formula" else _skip_reason_for_paragraph(raw_xml_elem)
            auto_num = _detect_auto_num(text, raw_xml_elem, backend.docx_obj)
            elements.append({
                "id": eid, "label": pending["label"], "level": pending["level"], "text": text,
                "loc": {"paraId": para_id} if para_id else None, "skip": skip, "auto_num": auto_num,
            })
            eid += 1
        pending = None

    for item, _level in doc.iterate_items():
        label = str(item.label)
        if label in FURNITURE_LABELS:
            continue
        if label == "table":
            flush()
            elements.append(_build_table_element(item, eid, docx_tables, table_index))
            table_index += 1
            eid += 1
            continue
        if label == "picture":
            flush()
            elements.append({"id": eid, "label": "picture", "level": None, "text": "",
                              "loc": None, "skip": "picture", "auto_num": None})
            eid += 1
            continue

        text = getattr(item, "text", "") or ""
        para_id = backend.obj_to_paraid.get(id(item))
        if para_id in table_paraids:  # 표 셀 문단 중복 방지(위 table_paraids 주석 참고)
            continue
        group_key = para_id or item.self_ref
        item_level = getattr(item, "level", None) if label == "section_header" else None

        if pending is not None and pending["group_key"] == group_key:
            if text:
                pending["text_parts"].append(text)
            continue
        flush()
        pending = {"group_key": group_key, "label": label, "level": item_level,
                   "para_id": para_id, "text_parts": [text] if text else []}
    flush()

    # 표 앞뒤 caption 연결(간단한 인접 검사, 6절 caption_ids)
    for i, e in enumerate(elements):
        if e["label"] != "table":
            continue
        for j in (i - 1, i + 1):
            if 0 <= j < len(elements) and elements[j]["label"] == "caption":
                e["caption_ids"].append(elements[j]["id"])

    loc_none = sum(1 for e in elements if e["label"] != "table" and e["loc"] is None)
    print(f"[docx parse] {Path(src_path).name}: 요소 {len(elements)}개, "
          f"loc=None(병합 위치 못 찾음, 청킹/검색/번역엔 영향 없음) {loc_none}개", flush=True)
    return elements


if __name__ == "__main__":
    from collections import Counter

    from docobj import add_spans, save_elements

    for f in sys.argv[1:]:
        elements = parse_docx(f)
        flat_text = add_spans(elements)
        counts = Counter(e["label"] for e in elements)
        skip_counts = Counter(e.get("skip") for e in elements if e.get("skip"))
        print(f"\n=== {Path(f).name} ({len(elements)}개 요소, flat_text {len(flat_text)}자) ===")
        print("label별 개수:", dict(counts))
        print("skip 사유별 개수:", dict(skip_counts))
        for e in elements[:10]:
            preview = e["text"][:40].replace("\n", "\\n")
            print(" ", e["id"], e["label"], e["level"], "|", e["loc"], "| skip=", e.get("skip"),
                  "| span=", e["span"], "|", preview)
        out = save_elements(Path(f).stem, "docx", elements, "data/processed")
        print("saved:", out)
