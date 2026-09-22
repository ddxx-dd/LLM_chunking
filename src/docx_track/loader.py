import re
import unicodedata
from pathlib import Path
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from langchain_core.documents import Document

SEP = "\n"

def normalize(s):
    s = unicodedata.normalize("NFC", s)
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub("[​‌‍‎‏﻿\xa0]", " ", s)
    s = re.sub(r"[ \t]+", " ", s)
    return s.strip()

class _Builder:
    """LangChain Document(page_content, metadata)로 조립 - Unit은 metadata["units"]의
    딕셔너리 리스트로 들어간다({"start":,"end":,"kind":,"meta":{...}})."""
    def __init__(self):
        self.parts = []
        self.units = []
        self.pos = 0

    def add(self, text, kind, meta=None):
        if not text:
            return
        start = self.pos
        end = self.pos + len(text)
        self.units.append({"start": start, "end": end, "kind": kind, "meta": meta or {}})
        self.parts.append(text + SEP)
        self.pos += len(text) + len(SEP)

    def done(self, name, fmt, log):
        return Document(page_content="".join(self.parts),
                         metadata={"name": name, "fmt": fmt, "log": log, "units": self.units})

def _table_rows_to_html(table):
    """python-docx Table -> HTML 문자열, gridSpan/vMerge 병합을 colspan/rowspan으로 보존.
    XML의 gridSpan/vMerge 속성을 직접 파싱하지 않는다 - row.cells가 병합된 그리드 위치마다
    "같은 셀 객체"를 반복해서 돌려주는 걸 병합 신호로 재활용한다(unstructured 라이브러리의
    기법을 이식, 직사각형 병합만 지원 - Word에서 실제로 만들어지는 병합은 거의 이 형태)."""
    matrix = []
    for row in table.rows:
        try:
            # id(cell._tc)를 키로 쓰면 안 된다 - 그 셀 객체 참조가 안 남아 GC되면 메모리
            # 주소(id)가 나중에 전혀 다른 셀에 재사용돼 우연히 같은 id가 나오는 충돌이
            # 실제로 발생함(실측 발견). 객체 자체를 키로 유지해 참조가 붙잡혀있게 한다.
            matrix.append([(cell.text.strip(), cell._tc) for cell in row.cells])
        except Exception:
            matrix.append([])
    if not matrix or not matrix[0]:
        return ""

    n_rows = len(matrix)
    n_cols = max(len(r) for r in matrix)
    for r in matrix:  # 행마다 열 수가 다른 기형적 표는 빈 칸으로 채워 정렬만 맞춤
        while len(r) < n_cols:
            r.append(("", object()))

    consumed = [[False] * n_cols for _ in range(n_rows)]
    trs = []
    for r in range(n_rows):
        tds = []
        for c in range(n_cols):
            if consumed[r][c]:
                continue
            consumed[r][c] = True
            text, key = matrix[r][c]

            colspan = 1  # 오른쪽으로 같은 key(=같은 셀 객체)가 이어지는 동안 확장
            while c + colspan < n_cols and not consumed[r][c + colspan] and matrix[r][c + colspan][1] == key:
                consumed[r][c + colspan] = True
                colspan += 1

            rowspan = 1  # 아래쪽으로 같은 폭 전체가 같은 key인 동안 확장
            while r + rowspan < n_rows and all(
                matrix[r + rowspan][c + dc][1] == key for dc in range(colspan)
            ):
                for dc in range(colspan):
                    consumed[r + rowspan][c + dc] = True
                rowspan += 1

            attrs = (f' colspan="{colspan}"' if colspan > 1 else "") + (f' rowspan="{rowspan}"' if rowspan > 1 else "")
            esc = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            tds.append(f"<td{attrs}>{esc}</td>" if esc else f"<td{attrs}/>")
        if tds:
            trs.append(f"<tr>{''.join(tds)}</tr>")
    return f"<table>{''.join(trs)}</table>"


def _iter_body(doc):
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)

def load_docx(filepath):
    doc = DocxDocument(filepath)
    b = _Builder()
    log = {"paras": 0, "tables": 0, "rows": 0}
    tbl_no = 0

    for item in _iter_body(doc):
        if isinstance(item, Paragraph):
            t = normalize(item.text)
            # 완전히 빈 문단만 버린다 - 1~2글자 소제목/라벨도 원본의 일부.
            if not t:
                continue
            style = item.style.name if item.style else ""
            is_head = style.startswith(("Heading", "제목", "Title"))
            b.add(t, "para", {"style": style, "heading": is_head})
            log["paras"] += 1

        elif isinstance(item, Table):
            tbl_no += 1
            log["tables"] += 1
            if not item.rows:
                continue
            # 표 전체를 마크업(HTML) 그대로 하나의 Unit으로 노출 - fixed/semantic 청커는
            # 이 태그 섞인 텍스트를 그냥 문자열로 취급해서 아무데서나 자를 수 있다(의도된
            # 정직한 약점). 구조를 아는 청커만 이 HTML을 파싱해서 행/열/병합을 복원해 쓴다.
            html_table = _table_rows_to_html(item)
            if html_table:
                b.add(html_table, "table", {"table": tbl_no, "rows": len(item.rows)})
                log["rows"] += len(item.rows)

    return b.done(Path(filepath).name, "docx", log)

def load_docx_bundle(paths):
    """여러 docx를 각각 독립된 Doc으로 로드(문서 경계를 넘어 청킹/검색하지 않도록)."""
    paths = sorted(Path(p) for p in paths)
    return [load_docx(str(p)) for p in paths]
