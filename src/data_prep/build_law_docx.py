"""국가법령정보센터(law.go.kr)에서 법령 전문을 받아 docx로 저장한다.
법령은 저작권법 제7조에 의해 애초에 저작권 보호 대상이 아니라(법령·고시·훈령 등),
라이선스 걱정 없이 쓸 수 있는 "대용량 단일 문서" 소스다.

law.go.kr의 조회 페이지(`/법령/<이름>`)는 실제 조문 내용이 없는 iframe 껍데기고,
진짜 조문은 그 iframe이 부르는 `LSW/lsInfoR.do?lsiSeq=...` 쪽에 있다(실측 확인) -
`편/장/절` 제목은 `<p class="gtit">`, 조문 하나하나는 `<p class="pty1_p4">`로
깔끔하게 구조화돼 있어서 그대로 파싱 가능하다.
"""
import re
import sys
import urllib.request
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument

sys.path.append(str(Path(__file__).resolve().parent.parent))
from config import DOCX_KOR_DIR

LAWS = {
    "민법": (284415, "20260317"),
    "형법": (284025, "20260913"),
    "상법": (273629, "20260910"),
    "민사소송법": (252393, "20250712"),
}

OUT_DIR = DOCX_KOR_DIR / "대용량문서"


def fetch_law_html(lsi_seq, ef_yd):
    url = (f"https://www.law.go.kr/LSW/lsInfoR.do?lsiSeq={lsi_seq}"
           f"&efYd={ef_yd}&chrClsCd=010202&urlMode=lsInfoR")
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as r:
        return r.read().decode("utf-8", errors="ignore")


def build_docx(name, lsi_seq, ef_yd, out_path):
    html = fetch_law_html(lsi_seq, ef_yd)
    soup = BeautifulSoup(html, "html.parser")

    doc = DocxDocument()
    doc.add_heading(name, level=1)

    n_heading, n_para = 0, 0
    for p in soup.select("p.gtit, p.pty1_p4"):
        text = re.sub(r"\s+", " ", p.get_text()).strip()
        if not text:
            continue
        if "gtit" in (p.get("class") or []):
            doc.add_heading(text, level=2)
            n_heading += 1
        else:
            doc.add_paragraph(text)
            n_para += 1

    doc.save(str(out_path))
    print(f"{name}: 편/장/절 제목 {n_heading}개 + 조문 {n_para}개 -> {out_path}")


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for name, (lsi_seq, ef_yd) in LAWS.items():
        build_docx(name, lsi_seq, ef_yd, OUT_DIR / f"{name}.docx")


if __name__ == "__main__":
    main()
