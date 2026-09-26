"""병합 전처리 점검(smart_chunk.md 8절 "병합" 행) - 표시 테스트(마킹) + 병합 지표.
docx 항등 테스트는 docx_track/identity_test.py, pdf 항등 테스트는 pdf_track/identity_test.py.
45개/50개 전체 스캔은 10절 3단계에서 이 모듈 함수들을 재사용한다."""
import sys
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.path.append(str(Path(__file__).resolve().parent))
from docx_track.parse import add_para_ids, parse_docx
from docx_track.writer import build_paraid_index, write_element_text

MARK_OPEN = "⟦"  # ⟦
MARK_CLOSE = "⟧"  # ⟧


def _in_fallback(p_xml):
    return any(a.tag.rsplit("}", 1)[-1] == "Fallback" for a in p_xml.iterancestors())


def _mark_targets(doc, loc):
    """이 loc이 실제로 걸쳐 있는 paraId들(설명/제외 용도) - table_index fallback도
    현재 문서 상태에서 셀 문단의 실제 paraId를 읽어 채운다(3단계 전체 스캔에서
    table_mismatch 표가 나와도 unwritten 오분류가 안 나게)."""
    if not loc:
        return []
    if "paraId" in loc:
        return [loc["paraId"]]
    if "paraIds" in loc:
        return list(loc["paraIds"])
    if "table_index" in loc:
        try:
            cell = doc.tables[loc["table_index"]].cell(loc["row"], loc["col"])
        except IndexError:
            return []
        return [p._p.get(qn("w14:paraId")) for p in cell.paragraphs]
    return []


def run_mark_test(src_path):
    """번역 가능한 모든 요소·셀에 ⟦원문⟧을 써서 저장한 뒤 다시 열어, 텍스트는 있는데
    ⟦가 없는 w:p를 찾는다. skip 사유로 설명되는 건 정상, 안 되는 건 "unwritten"으로 분류.
    write_element_text가 기본으로 mc:AlternateContent의 Fallback 쪽 대응 문단에도 같은
    텍스트를 쓰므로(docx_track/writer.py의 fallback_sibling, ★ 실측 확인된 버그의 해결책),
    Fallback 문단은 정상적으로 마킹돼 있어야 한다."""
    src_path = Path(src_path)
    anchored_path = src_path.with_suffix(".anchored.docx")
    if not anchored_path.exists():
        anchored_path = add_para_ids(src_path)
    elements = parse_docx(src_path)

    doc = DocxDocument(str(anchored_path))
    paraid_to_p = build_paraid_index(doc)

    paraid_reason = {}  # paraId -> skip 사유(설명 가능한 것들) 또는 "picture"
    marked = 0

    def handle(loc, text, skip_reason, auto_num=None):
        nonlocal marked
        pids = _mark_targets(doc, loc)
        for pid in pids:
            if pid:
                paraid_reason[pid] = skip_reason
        if skip_reason:
            return
        result = write_element_text(doc, paraid_to_p, loc, f"{MARK_OPEN}{text}{MARK_CLOSE}", auto_num=auto_num)
        if result == "ok":
            marked += 1

    for e in elements:
        if e["label"] == "table":
            for c in e["cells"]:
                handle(c.get("loc"), c["text"], c.get("skip") or e.get("skip"))
            continue
        reason = "picture" if e["label"] == "picture" else e.get("skip")
        handle(e.get("loc"), e["text"], reason, e.get("auto_num"))

    out_path = src_path.with_name(src_path.stem + ".marktest.docx")
    doc.save(str(out_path))

    doc2 = DocxDocument(str(out_path))
    unwritten = []
    explained = 0
    fallback_unwritten = 0
    for p in doc2.element.body.iter(qn("w:p")):
        text = Paragraph(p, doc2).text
        if not text.strip():
            continue
        if MARK_OPEN in text:
            continue
        pid = p.get(qn("w14:paraId"))
        reason = paraid_reason.get(pid)
        if reason:
            explained += 1
            continue
        is_fb = _in_fallback(p)
        if is_fb:
            fallback_unwritten += 1
        unwritten.append({"paraId": pid, "fallback": is_fb, "text": text[:50]})

    print(f"[mark test] {src_path.name}")
    print(f"  마킹 성공: {marked}, skip 사유로 설명됨(정상): {explained}")
    print(f"  unwritten(설명 안 됨): {len(unwritten)} (그중 Fallback: {fallback_unwritten})")
    for u in unwritten[:10]:
        print(f"    paraId={u['paraId']} fallback={u['fallback']} text={u['text']!r}")
    return {"marked": marked, "explained": explained, "unwritten": len(unwritten),
            "unwritten_fallback": fallback_unwritten, "examples": unwritten[:10]}


def count_mixed_formatting_paragraphs(elements, anchored_path):
    """병합 지표: run마다 bold/italic 조합이 다른 문단 수(알려진 한계 - writer가 첫 w:t에
    번역문 전체를 넣으므로, 문단 안에 서식이 섞여 있으면 번역 후 서식이 첫 run 쪽으로
    쏠린다. 대상은 실제로 번역 대상이 되는 요소/셀의 paraId만 센다)."""
    doc = DocxDocument(str(anchored_path))
    paraid_to_p = build_paraid_index(doc)

    target_paraids = set()
    for e in elements:
        if e["label"] == "table":
            for c in e["cells"]:
                if c.get("skip") or e.get("skip"):
                    continue
                loc = c.get("loc") or {}
                target_paraids.update(loc.get("paraIds", []))
            continue
        if e.get("skip") or e["label"] == "picture":
            continue
        loc = e.get("loc") or {}
        if "paraId" in loc:
            target_paraids.add(loc["paraId"])

    mixed = 0
    for pid in target_paraids:
        p = paraid_to_p.get(pid)
        if p is None:
            continue
        para = Paragraph(p, doc)
        combos = {(bool(r.bold), bool(r.italic)) for r in para.runs if r.text}
        if len(combos) > 1:
            mixed += 1
    return {"mixed_formatting_paragraphs": mixed, "target_paragraphs": len(target_paraids)}


if __name__ == "__main__":
    for f in sys.argv[1:]:
        result = run_mark_test(f)
        anchored = Path(f).with_suffix(".anchored.docx")
        elements = parse_docx(f)
        metric = count_mixed_formatting_paragraphs(elements, anchored)
        print(f"  서식 섞인 문단 수(알려진 한계 - 번역 후 서식이 첫 run 쪽으로 쏠림): "
              f"{metric['mixed_formatting_paragraphs']} / {metric['target_paragraphs']}")
