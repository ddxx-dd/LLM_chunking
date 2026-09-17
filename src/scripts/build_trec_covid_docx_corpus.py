"""BeIR/trec-covid(corpus, cc-by-sa-4.0) + BeIR/trec-covid-qrels(TREC 판정) -> DOCX 검색 코퍼스.

TREC-COVID는 이번에 확정한 4개 소스 중 유일하게 등급화된(graded) 관련도(-1/0/1/2)를 갖는
소스다. 쿼리 8개를 뽑아 각 쿼리마다 score=2(고관련)/1(관련)/0(무관) 문서를 최대 3개씩
고르고(쿼리 간 중복 문서는 자동 제거), qa.json에 문서별 등급을 그대로 남겨 nDCG@k
계산이 가능하게 한다.

문서(corpus)가 논문 전문이 아니라 초록급이라 짧은 경우가 많다(중앙값 1,361자). 무작위
대신 각 등급 버킷 안에서 "가장 긴 문서"부터 고른다 - corpus 전체엔 122,392자짜리 문서도
있어서(전문이 포함된 항목도 섞여있음) 무작위보다 훨씬 길어진다.
"""
import html
import json
import random
import sys
from pathlib import Path

import pandas as pd
from docx import Document as DocxDocument
from huggingface_hub import hf_hub_download

sys.path.append(str(Path(__file__).resolve().parent.parent))
from config import DOCX_ENG_DIR

N_QUERIES = 8
PER_SCORE_CAP = 3
SEED = 42

OUT_DIR = DOCX_ENG_DIR / "trec_covid"
DOCS_DIR = OUT_DIR / "docs"


def write_doc(path, title, text):
    d = DocxDocument()
    d.add_heading(html.unescape(title) or "(no title)", level=1)
    for para in text.split("\n"):
        # CORD-19 원문 일부에 &gt; 같은 미해제 HTML 엔티티가 그대로 섞여있음(우리
        # 코드가 만든 게 아니라 원본 trec-covid corpus 자체의 특징) - 그대로 두면
        # Word에서 "&gt;"라는 글자가 그대로 보임.
        para = html.unescape(para.strip())
        if para:
            d.add_paragraph(para)
    d.save(str(path))


def main():
    random.seed(SEED)
    DOCS_DIR.mkdir(parents=True, exist_ok=True)

    queries = pd.read_parquet(hf_hub_download("BeIR/trec-covid", repo_type="dataset", filename="queries/queries-00000-of-00001.parquet"))
    qrels = pd.read_csv(hf_hub_download("BeIR/trec-covid-qrels", repo_type="dataset", filename="test.tsv"), sep="\t")
    print(f"queries={len(queries)} qrels={len(qrels)}")

    corpus_full = pd.read_parquet(hf_hub_download("BeIR/trec-covid", repo_type="dataset", filename="corpus/corpus-00000-of-00001.parquet"))
    doc_len = corpus_full.set_index("_id")["text"].str.len()

    picked_qids = random.sample(queries["_id"].astype(str).tolist(), N_QUERIES)
    qrels["query-id"] = qrels["query-id"].astype(str)

    per_query_docs = {}
    needed_ids = set()
    for qid in picked_qids:
        sub = qrels[qrels["query-id"] == qid]
        chosen = []
        for score in (2, 1, 0):
            pool = sub[sub["score"] == score]["corpus-id"]
            pool = pool[pool.isin(doc_len.index)]
            top = doc_len.loc[pool].sort_values(ascending=False).head(PER_SCORE_CAP)
            chosen.extend((cid, score) for cid in top.index)
        per_query_docs[qid] = chosen
        needed_ids.update(cid for cid, _ in chosen)
    print(f"필요 문서(중복 제거): {len(needed_ids)}개")

    corpus = corpus_full[corpus_full["_id"].isin(needed_ids)].set_index("_id")
    print(f"코퍼스에서 실제 발견된 문서: {len(corpus)}개")

    id_to_filename = {}
    for i, cid in enumerate(corpus.index, 1):
        row = corpus.loc[cid]
        title = row["title"] if isinstance(row["title"], str) and row["title"].strip() else f"doc_{cid}"
        safe = "".join(c for c in title if c not in '\\/:*?"<>|').strip()[:60]
        fname = f"{cid}_{safe}.docx"
        id_to_filename[cid] = fname
        write_doc(DOCS_DIR / fname, title, row["text"])
        print(f"[{i}/{len(corpus)}] {fname}")

    qtext = queries.set_index(queries["_id"].astype(str))["text"]
    qa = []
    for qid in picked_qids:
        docs = [
            {"doc_id": cid, "filename": id_to_filename[cid], "score": score}
            for cid, score in per_query_docs[qid] if cid in id_to_filename
        ]
        qa.append({"query_id": qid, "query": qtext.loc[qid], "graded_docs": docs})

    (OUT_DIR / "qa.json").write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"완료: 문서 {len(corpus)}개, 쿼리 {len(qa)}개")
    print(f"저장 위치: {OUT_DIR}")


if __name__ == "__main__":
    main()
