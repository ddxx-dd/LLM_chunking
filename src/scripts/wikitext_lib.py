"""MediaWiki(위키책/위키문헌/영어 위키책 등) wikitext -> docx 공용 변환 유틸.

`Special:Export`로 받은 raw wikitext를 파싱해서 헤딩(mwparserfromhell)과 표(`{|...|}`,
mwparserfromhell이 지원 안 해서 직접 정규식 파싱)를 순서대로 뽑아 docx로 쓴다.
"""
import re
import urllib.parse
import urllib.request

import mwparserfromhell as mwph
from docx import Document as DocxDocument

# 닫는 줄이 "|}"가 아니라 " |}"처럼 앞에 공백이 붙는 경우가 실제로 있어(예: 위키책
# "인코딩과 문자 집합/Ascii85") \n 뒤에 공백을 허용해야 함 - 안 그러면 이 표가 통째로
# 안 걸러지고 원문 그대로 문단에 새어나감.
#
# 표 안에 표가 또 들어있는 경우가 있음(예: 위키백과 "한글"의 자모표 - 닿소리/홀소리 표를
# 나란히 배치하려고 바깥에 레이아웃용 표로 한 번 더 감쌈). non-greedy라 바깥 "{|"부터
# 시작해서 제일 먼저 만나는 "|}"(안쪽 표의 닫는 줄)에서 끊어버리면 바깥/안쪽 표가 뒤섞임 -
# "(?!\{\|)."로 자기 안에 또 다른 "{|"가 없는, 즉 가장 안쪽 표만 매칭하도록 한다. 이걸
# 반복 적용하면(바깥 표는 첫 라운드에 안쪽 표가 플레이스홀더로 바뀌어 다음 라운드에
# "중첩 없음" 상태가 됨) 몇 겹이 중첩됐어도 안쪽부터 차례로 다 잡힌다.
TABLE_BLOCK_RE = re.compile(r"\{\|(?:(?!\{\|).)*?\n[ \t]*\|\}", re.S)
WIKILINK_STRIP_RE = re.compile(r"\[\[(?:[^\]|]*\|)?([^\]]*)\]\]")


def export_pages(site, titles):
    """site 예: 'ko.wikibooks.org'. titles: 페이지 제목 리스트. 반환: {title: wikitext}."""
    url = f"https://{site}/wiki/Special:Export"
    data = urllib.parse.urlencode({"pages": "\n".join(titles), "curonly": "1"}).encode()
    req = urllib.request.Request(url, data=data, headers={"User-Agent": "Mozilla/5.0"})
    xml = urllib.request.urlopen(req, timeout=30).read().decode("utf-8")
    pages = re.findall(r"<title>(.*?)</title>.*?<text[^>]*>(.*?)</text>", xml, re.S)
    out = {}
    for title, text in pages:
        out[_unescape(title)] = _unescape(text)
    return out


def _unescape(s):
    return (s.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
             .replace("&quot;", '"').replace("&#039;", "'"))


def _strip_file_embeds(text):
    """[[파일:...]]/[[File:...]]/[[Image:...]] 통째로 제거(캡션에 실제 프로즈가 아니라
    이미지 설명이 들어있고, 캡션 안에 또 [[...]] 링크가 중첩되는 경우가 많아 일반
    WIKILINK_STRIP_RE(정규식이라 중첩 괄호를 못 다룸)로는 못 건너뛰고 깨진 채 새어나감 -
    예: 천체물리학/외계행성에서 "FFJFRULKTORIR" 잔재. 괄호 깊이를 직접 세서 제거."""
    out = []
    i, n = 0, len(text)
    while i < n:
        if text[i:i + 2] == "[[" and re.match(r"\[\[(파일|File|Image):", text[i:i + 20]):
            depth = 1
            j = i + 2
            while j < n and depth > 0:
                if text[j:j + 2] == "[[":
                    depth += 1
                    j += 2
                elif text[j:j + 2] == "]]":
                    depth -= 1
                    j += 2
                else:
                    j += 1
            i = j
        else:
            out.append(text[i])
            i += 1
    return "".join(out)


def _clean_inline(text):
    text = _strip_file_embeds(text)
    for _ in range(3):  # 남은 [[...]]가 중첩돼 있으면 안쪽부터 여러 번 돌려야 다 풀림
        new_text = WIKILINK_STRIP_RE.sub(r"\1", text)
        if new_text == text:
            break
        text = new_text
    text = re.sub(r"'''''(.*?)'''''", r"\1", text)
    text = re.sub(r"'''(.*?)'''", r"\1", text)
    text = re.sub(r"''(.*?)''", r"\1", text)
    text = re.sub(r"\[https?://\S+\s+([^\]]*)\]", r"\1", text)
    text = re.sub(r"<ref[^>]*>.*?</ref>", "", text, flags=re.S)
    text = re.sub(r"<ref[^>]*/>", "", text)
    # <math>...</math>는 렌더링 안 하면 LaTeX 소스 코드 그대로 문단에 남아 문맥이
    # 끊김(사용자가 위키백과 "일반 상대성이론"에서 직접 발견 - 142곳). <gallery>...
    # </gallery>는 안에 "Image:파일.jpg|캡션" 위키 문법 줄이 그대로 들어있어 역시
    # 프로즈가 아님(영어 위키책 "World History/Maps"에서 발견). 둘 다 렌더링 불가능한
    # 비-프로즈 콘텐츠라 통째로 제거한다(<ref>와 같은 처리).
    text = re.sub(r"<math[^>]*>.*?</math>", "", text, flags=re.S)
    text = re.sub(r"<gallery[^>]*>.*?</gallery>", "", text, flags=re.S)
    text = re.sub(r"<[^>]+>", "", text)
    # 소스 자체가 깨져서(닫는 |}/''' 누락 등 - 위키책 편집 오류, 우리 쪽 문제 아님) 정상
    # 처리가 안 되는 극소수 표/서식 잔재를 방어적으로 제거(내용은 보존, 문법 기호만 제거)
    text = re.sub(r"^\{\|.*$", "", text, flags=re.M)
    text = re.sub(r"^\s*\|[-}]\s*$", "", text, flags=re.M)
    text = re.sub(r"^!\s*", "", text, flags=re.M)
    text = re.sub(r"^\|\s*", "", text, flags=re.M)
    text = re.sub(r"'{2,}", "", text)
    # 위키 목록 마커(* 글머리표, # 번호목록, ; 정의어, : 정의)도 mwparserfromhell이
    # 특수 처리 안 해서 그냥 문자 그대로 남음 - 줄 앞의 마커만 제거(내용은 유지).
    text = re.sub(r"^[*#;:]+\s*", "", text, flags=re.M)
    return text.strip()


def _strip_cell_attrs(cell):
    """위키 표 셀은 "속성들 |실제내용" 형태로 셀 하나 안에 파이프가 한 번 더 나올 수
    있음(예: align="center" bgcolor="#eaf6ea" |첫소리). 속성부에는 항상 "="가 있으므로
    그걸로 판별해서 마지막 "|" 뒤의 실제 내용만 남긴다 - 안 그러면 "align=..." 잡음이
    그대로 셀 내용에 섞여나감(사용자가 위키백과 "한글" 문서에서 직접 발견)."""
    cell = cell.strip()
    if "|" in cell:
        attr_part, content = cell.rsplit("|", 1)
        if "=" in attr_part:
            return content
    return cell


def _parse_table(block):
    """{|...|} 블록 -> [[cell,cell,...], ...] (첫 행이 헤더). 못 파싱하면 None.

    colspan/rowspan을 쓰는 표는 건너뛴다 - 이건 실제로 "표"가 아니라 시각적 배치를
    위한 그림(예: 위키백과 "한글"의 초성/중성/종성 위치를 보여주는 자모 조립 다이어그램)
    인 경우가 많고, 우리 파서는 셀 병합을 이해 못 해서 그대로 평평한 표로 펴면 같은
    글자가 중복되거나 빈 칸이 생기는 등 원문과 다른 엉뚱한 모양이 됨(사용자가 직접
    발견). 이런 다이어그램이 설명하는 규칙은 대개 주변 문단에 말로도 이미 나와 있어서,
    억지로 잘못 재구성하는 것보다 건너뛰는 게 원문에 더 충실하다."""
    if re.search(r"(?:colspan|rowspan)\s*=", block):
        return None
    rows = []
    cur = []
    for line in block.split("\n"):
        line = line.strip()
        if line.startswith("{|") or line.startswith("|}"):
            continue
        if line.startswith("|+"):
            # 표 캡션(|+ 표제목) 줄 - 데이터 행이 아니라 건너뛴다. 안 그러면
            # "+ 닿소리"처럼 캡션이 첫 행의 셀 하나로 잘못 섞여 들어감.
            continue
        if line.startswith("|-"):
            if cur:
                rows.append(cur)
                cur = []
            continue
        if line.startswith("!"):
            cells = re.split(r"!!|\|\|", line[1:])
            cur.extend(_clean_inline(_strip_cell_attrs(c)) for c in cells)
        elif line.startswith("|"):
            cells = re.split(r"\|\|", line[1:])
            cur.extend(_clean_inline(_strip_cell_attrs(c)) for c in cells)
    if cur:
        rows.append(cur)
    rows = [r for r in rows if any(c.strip() for c in r)]
    if len(rows) < 2:
        return None
    ncols = max(len(r) for r in rows)
    if ncols < 2:
        # 열이 1개뿐이면 "표"라고 부를 실질적 구조(헤더 vs 데이터, 행 vs 열)가 없음 -
        # 워드로 열면 그냥 단어 하나씩 세로로 늘어선 게 문단처럼 보임(사용자가 지적한
        # 증상과 정확히 일치) - 대개 색깔로 꾸민 다이어그램의 잔재라 표로 렌더링 안 함.
        return None
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    return rows


def _resolve_table_events(key, placeholders, seen=None):
    """플레이스홀더 하나를 실제 ("table", rows) 이벤트 리스트로 푼다. 안에 또 다른
    플레이스홀더가 들어있으면(= 레이아웃용 바깥 표, 실제 데이터가 아님) 그 표 자체는
    버리고 안에 들어있는 진짜 표들만 순서대로 재귀적으로 풀어낸다."""
    if seen is None:
        seen = set()
    if key in seen:  # 방어적 - 순환 참조는 있을 수 없지만 만일을 위해
        return []
    seen.add(key)
    raw = placeholders[key]
    nested = list(NESTED_PLACEHOLDER_RE.finditer(raw))
    if nested:
        out = []
        for m in nested:
            out.extend(_resolve_table_events(f"@@TABLE{m.group(1)}@@", placeholders, seen))
        return out
    rows = _parse_table(raw)
    return [("table", rows)] if rows else []


NESTED_PLACEHOLDER_RE = re.compile(r"@@TABLE(\d+)@@")


def wikitext_events(wikitext):
    """wikitext 한 페이지 -> [("heading", level, text) | ("para", text) | ("table", rows), ...] 순서대로."""
    placeholders = {}

    def _stash(m):
        key = f"@@TABLE{len(placeholders)}@@"
        placeholders[key] = m.group(0)
        return "\n" + key + "\n"

    # 한 번에 다 못 잡음 - 표 안에 표가 중첩돼 있으면 안쪽 표가 플레이스홀더로 바뀐
    # 뒤에야 바깥 표도 "중첩 없음" 상태가 되어 다음 라운드에 잡힌다. 더 이상 안 바뀔
    # 때까지 반복.
    stashed = wikitext
    while True:
        new_stashed = TABLE_BLOCK_RE.sub(_stash, stashed)
        if new_stashed == stashed:
            break
        stashed = new_stashed

    code = mwph.parse(stashed)
    events = []
    buf = []

    def flush():
        # buf의 각 원소는 mwparserfromhell이 쪼갠 노드 하나하나(문장 중간의 <sub>2</sub>,
        # [[위키링크]] 등도 전부 별도 노드)라서 "\n"으로 이어붙이면 "CO\n2\n + 2H\n2\nO"처럼
        # 문장 중간이 줄바꿈투성이가 됨(사용자가 직접 발견) - 그냥 이어붙인다. 실제 줄바꿈은
        # 노드 문자열 자체에 이미 원본 그대로 들어있으므로 따로 넣을 필요 없음.
        text = _clean_inline("".join(buf)).strip()
        buf.clear()
        if text:
            for para in re.split(r"\n{2,}", text):
                para = para.strip()
                if para and not para.startswith(("{{", "|", "[[분류", "[[Category")):
                    events.append(("para", para))

    for node in code.nodes:
        if isinstance(node, mwph.nodes.Heading):
            flush()
            events.append(("heading", node.level, _clean_inline(str(node.title))))
        elif isinstance(node, mwph.nodes.Template):
            # {{llang|en|photosynthesis}}, {{돌아가기|...}} 같은 위키 템플릿은 그대로 두면
            # 원본에 없던 "{{...}}" 잡음으로 보임 - 템플릿 확장 기능은 없으니 통째로 버린다
            # (본문 프로즈에 꼭 필요한 정보가 아니라 부가 주석/내비게이션인 경우가 대부분).
            continue
        elif isinstance(node, mwph.nodes.Wikilink) and re.match(r"^(분류|Category|category):|^[a-z]{2,3}:\S", str(node.title)):
            # [[분류:한국어]](카테고리)나 [[de:Koreanisch]](다국어판 링크)는 렌더링 시
            # 사이드바로 빠지지 실제 본문에 절대 안 나온다 - 원문에도 "본문"이 아니므로
            # 이걸 문단으로 남기는 게 오히려 "원본과 다르게" 보이는 것 - 건너뛴다.
            continue
        else:
            s = str(node)
            # 플레이스홀더가 주변 문단 텍스트와 같은 노드에 합쳐지는 경우가 있어(표 앞뒤에
            # 위키 문법 경계가 없으면 mwparserfromhell이 통째로 하나의 텍스트 노드로 묶음),
            # 노드 전체가 플레이스홀더인지만 보면 못 잡는다 - 노드 안에서 검색해서 잘라낸다.
            pos = 0
            for m in NESTED_PLACEHOLDER_RE.finditer(s):
                buf.append(s[pos:m.start()])
                flush()
                key = f"@@TABLE{m.group(1)}@@"
                events.extend(_resolve_table_events(key, placeholders))
                pos = m.end()
            buf.append(s[pos:])
    flush()
    return _drop_empty_headings(events)


def _drop_empty_headings(events):
    """소제목만 있고 그 아래(같거나 더 상위 레벨 헤딩이 나오기 전까지) 문단/표가 하나도
    없는 헤딩은 버린다. 실제 위키 원문에 편집자가 절 제목만 만들어놓고 아직 내용을 안 쓴
    "스텁 섹션"이 흔함(예: 영어 위키책 World History의 여러 챕터 - "Authoritarian
    governments", "Italy", "The Gupta Empire" 등이 전부 "{{BookCat}}"만 있고 본문이
    없는 채로 라이브 사이트에 실제로 그렇게 올라가 있음, 직접 원문 대조로 확인). 내용이
    아예 없으면 정보량이 0이라 문서 끝에 뜬 소제목처럼 보여서 오류로 오인되기 쉬움 -
    제거한다. 자기 하위에 실제 내용이 있는 상위 헤딩은 안 지운다(예: H2 다음에 H3가
    오고 H3 밑에 진짜 문단이 있으면 H2도 H3도 둘 다 유지)."""
    n = len(events)
    keep = [True] * n
    for i, ev in enumerate(events):
        if ev[0] != "heading":
            continue
        level = ev[1]
        has_content = False
        for j in range(i + 1, n):
            ev2 = events[j]
            if ev2[0] == "heading" and ev2[1] <= level:
                break
            if ev2[0] in ("para", "table"):
                has_content = True
                break
        keep[i] = has_content
    return [ev for ev, k in zip(events, keep) if k]


def events_to_docx(events, title, out_path, base_level=1):
    d = DocxDocument()
    d.add_heading(title, level=base_level)
    for ev in events:
        if ev[0] == "heading":
            level, text = ev[1], ev[2]
            if text:
                d.add_heading(text, level=min(base_level + level, 9))
        elif ev[0] == "para":
            d.add_paragraph(ev[1])
        elif ev[0] == "table":
            rows = ev[1]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    t.cell(ri, ci).text = val
    d.save(str(out_path))


def page_to_docx(site, title, out_path, display_title=None):
    pages = export_pages(site, [title])
    wikitext = pages.get(title, "")
    events = wikitext_events(wikitext)
    events_to_docx(events, display_title or title, out_path)
    return events


def list_book_pages(site, prefix):
    """prefix(책 제목)로 시작하는 실제 하위 페이지 전부(루트 포함) 제목 리스트."""
    url = f"https://{site}/w/api.php?" + urllib.parse.urlencode({
        "action": "query", "list": "allpages", "apprefix": prefix,
        "apnamespace": "0", "aplimit": "200", "format": "json",
    })
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    import json
    d = json.loads(urllib.request.urlopen(req, timeout=20).read())
    return [p["title"] for p in d["query"]["allpages"]]


TOC_LINK_RE = re.compile(r"^[#*]+\s*\[\[([^\]|]+)")


def _toc_order(root_wikitext):
    """루트 페이지의 번호목록(#)이나 글머리표(*)로 된 목차에서 하위 페이지의 진짜
    읽기 순서를 뽑는다. list_book_pages()가 쓰는 MediaWiki API는 하위 페이지를
    알파벳순으로만 주기 때문에(예: "일본어 입문"에서 "50음도"가 "1단 동사"보다 알파벳상
    뒤에 와서, 가나부터 배우는 게 정상인데 동사 활용을 먼저 배우는 순서로 나옴 - 사용자가
    지적해서 발견), 진짜 순서는 루트 페이지 자신의 목차에서 가져와야 한다. 책마다 목차
    형식이 다름(번호목록/글머리표/중첩 등) - 전부 "#"나 "*"로 시작하는 줄의 첫 위키링크를
    순서대로 모으면 다 잡힌다(체코어처럼 목차 자체가 없는 참고문법형 책은 빈 리스트 반환 -
    이 경우는 알파벳 순서를 그대로 쓴다, 애초에 선형적 순서가 없는 책이라 무방)."""
    order = []
    seen = set()
    for line in root_wikitext.split("\n"):
        m = TOC_LINK_RE.match(line.strip())
        if m:
            target = m.group(1).strip()
            if target not in seen:
                seen.add(target)
                order.append(target)
    return order


def book_to_docx(site, prefix, out_path, display_title=None, page_titles=None):
    """책 전체(루트+하위페이지)를 페이지별 Heading 2로 구분해 docx 1개로 합친다.
    순서는 루트 페이지의 목차(_toc_order)를 따르고, 목차에 없는 페이지(고아/오타 페이지 등)
    는 원래(API 알파벳) 순서 그대로 뒤에 붙인다 - 실제 내용은 하나도 빠뜨리지 않음."""
    titles = page_titles or list_book_pages(site, prefix)
    pages = export_pages(site, titles)
    toc = _toc_order(pages.get(prefix, ""))
    toc_index = {t: i for i, t in enumerate(toc)}
    orig_index = {t: i for i, t in enumerate(titles)}

    def sort_key(t):
        if t == prefix:
            return (-1, 0)
        if t in toc_index:
            return (0, toc_index[t])
        return (1, orig_index[t])

    titles = sorted(titles, key=sort_key)
    d = DocxDocument()
    d.add_heading(display_title or prefix, level=1)
    total_chars = 0
    for t in titles:
        wikitext = pages.get(t)
        if not wikitext:
            continue
        events = wikitext_events(wikitext)
        if not events:
            continue
        section_title = t[len(prefix):].lstrip("/") or t
        d.add_heading(section_title, level=2)
        for ev in events:
            if ev[0] == "heading":
                level, text = ev[1], ev[2]
                if text:
                    d.add_heading(text, level=min(2 + level, 9))
                    total_chars += len(text)
            elif ev[0] == "para":
                d.add_paragraph(ev[1])
                total_chars += len(ev[1])
            elif ev[0] == "table":
                rows = ev[1]
                tb = d.add_table(rows=len(rows), cols=len(rows[0]))
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row):
                        tb.cell(ri, ci).text = val
                        total_chars += len(val)
    d.save(str(out_path))
    return len(titles), total_chars
