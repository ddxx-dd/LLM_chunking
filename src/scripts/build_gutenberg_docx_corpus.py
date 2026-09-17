"""Project Gutenberg(퍼블릭 도메인) 고전문학 -> DOCX, 챕터별로 분리.

Pride and Prejudice(id 1342) 사용. 책 전체(72.8만자)를 한 파일로 두면 다른 문서
대비 압도적으로 커지므로, 실제 "CHAPTER" 마커 기준으로 챕터 단위 docx로 쪼갠다
(위키책/위키문헌 장편소설을 챕터 단위로 다룬 것과 같은 이유). Gutenberg 머리말/
꼬리말(라이선스 고지문 등)은 START/END 마커로 제거.

원본 텍스트 자체에 두 가지 표기 불일치가 있어(사용자가 직접 챕터 순서를 확인해달라고
해서 발견) 처음엔 챕터 1개(1장)가 통째로 빠지고 2개(13,14장)도 빠졌었다:
1. **1장은 "CHAPTER I."처럼 독립된 줄이 아니라 삽화 캡션 안에 "Chapter I.]"로 박혀있음**
   (이 판본의 장식체 1장 표지 페이지 특징 - 다른 장은 전부 평범한 "CHAPTER N." 줄임).
2. **13장·14장만 표기가 "CHAPTER XIII"/"CHAPTER XIV"로 마침표가 없음**(원문 자체의
   오타/조판 불일치, 나머지 59개 장은 다 마침표 있음).
정규식을 대소문자·마침표 유무·삽화 대괄호까지 허용하도록 넓혀서 61장 전부(빠짐/중복
없음, 직접 로마숫자 연속성 검증함) 잡히게 했다.
"""
import re
import sys
import urllib.request
from pathlib import Path

from docx import Document as DocxDocument

sys.path.append(str(Path(__file__).resolve().parent.parent))
from config import DOCX_ENG_DIR

OUT_DIR = DOCX_ENG_DIR / "Gutenberg_PrideAndPrejudice"
URL = "https://www.gutenberg.org/cache/epub/1342/pg1342.txt"

CHAPTER_RE = re.compile(r"^(?:CHAPTER|Chapter)\s+([IVXLCDM]+)\.?\]?\s*$", re.M)


def _strip_illustrations(text):
    """"[Illustration: 캡션 [_Copyright 1894 by George Allen._]]"처럼 대괄호가
    중첩돼 있어(원본 1894년 삽화판 표기) 단순 비탐욕 정규식으로는 안쪽 "]"에서 잘못
    끊긴다 - 괄호 깊이를 직접 세어 통째로 제거한다(전체 2371개 문단 중 155개, 6.5%가
    이 잔재였음). 삽화 설명/저작권 표기는 본문 프로즈가 아니라 제거해도 내용 손실 없음."""
    out = []
    i, n = 0, len(text)
    while i < n:
        if text[i:i + 13] == "[Illustration":
            depth = 1
            j = i + 1
            while j < n and depth > 0:
                if text[j] == "[":
                    depth += 1
                elif text[j] == "]":
                    depth -= 1
                j += 1
            i = j
        else:
            out.append(text[i])
            i += 1
    return "".join(out)


def fetch_text():
    # 여기서 삽화를 지우면 안 됨 - 1장의 챕터 마커("Chapter I.")가 바로 이 삽화 대괄호
    # 안에 같이 박혀있어서(예: "[Illustration: ...\n\nChapter I.]"), 먼저 지워버리면
    # 챕터 경계 자체가 사라짐. 챕터로 다 쪼갠 뒤(각 챕터 본문은 마커 뒤부터 시작이라
    # 마커와 안 겹침) split_chapters()에서 각 본문에 대해서만 지운다.
    req = urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0"})
    raw = urllib.request.urlopen(req, timeout=30).read().decode("utf-8")
    start = raw.find("*** START OF THE PROJECT GUTENBERG EBOOK")
    end = raw.find("*** END OF THE PROJECT GUTENBERG EBOOK")
    start = raw.find("\n", start) + 1
    return raw[start:end].strip()


def split_chapters(text):
    matches = list(CHAPTER_RE.finditer(text))
    chapters = []
    for i, m in enumerate(matches):
        body_start = m.end()
        body_end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        body = _strip_illustrations(text[body_start:body_end].strip())
        chapters.append((m.group(1), body))
    return chapters


def write_doc(chap_num, body, out_path):
    d = DocxDocument()
    d.add_heading(f"Pride and Prejudice - Chapter {chap_num}", level=1)
    for para in re.split(r"\n\s*\n", body):
        para = " ".join(para.split())
        if para:
            d.add_paragraph(para)
    d.save(str(out_path))


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    text = fetch_text()
    chapters = split_chapters(text)
    print(f"챕터 {len(chapters)}개 발견")
    for i, (num, body) in enumerate(chapters, 1):
        out = OUT_DIR / f"chapter_{i:02d}.docx"
        write_doc(num, body, out)
        print(f"[{i}/{len(chapters)}] Chapter {num}: {len(body)}자")
    print(f"완료: {len(chapters)}개")
    print(f"저장 위치: {OUT_DIR}")


if __name__ == "__main__":
    main()
