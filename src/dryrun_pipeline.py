"""10절 3단계: pipeline.py 전체 dry-run(--fake-llm) - docx 45개·pdf 50개·srt 전체 ×
fixed/semantic/smart 세 청커. 문서/영화당 파싱·group_elements는 1번만(캐시 재사용,
세 청커가 같은 N에 맞춰지므로 group_elements를 청커마다 다시 돌리면 3배 낭비 -
pipeline.py의 build_chunks는 단일 청커 호출용이라 그대로 안 쓰고 여기서만 한 번에
셋을 같이 만든다) 돌려서 설계서 10절 3번 체크리스트를 전부 실측으로 채운다."""
import json
import re
import sys
import time
import traceback
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent))

from config import RESULTS_DIR, SRT_ENG_DIR, SUBTITLE_DATASETS
from docobj import add_spans, cached_parse, get_target, is_skip_owner, owner_fragments_for_span
from pipeline import (PIPELINE_RESULTS_DIR, PROMPT_TEMPLATES, docx_files, fake_llm, merge_docx,
                       merge_pdf, parse_numbered, pdf_files, setup, translate_elements)
from smart_chunker import group_elements, make_fixed_chunks, make_semantic_chunks, section_for_owners, section_for_span

DRYRUN_DIR = RESULTS_DIR / "checks"
DRYRUN_DIR.mkdir(parents=True, exist_ok=True)
MARK_RE = re.compile(r"[«»]")


def _fragment_chunk(elements, flat_text, s, e):
    frags = owner_fragments_for_span(elements, s, e)
    return {"owners": [o for o, _ in frags], "fragments": frags, "text": flat_text[s:e],
            "section": section_for_span(elements, s)}


def build_all_chunks(elements, flat_text, embeddings, count_tokens, min_tokens=100, max_tokens=400, window=2):
    """세 청커를 group_elements 1번으로 같이 만든다(문서당 3배 낭비 방지) - 반환:
    {"smart": (chunks, extra), "fixed": (chunks, extra), "semantic": (chunks, extra)}.
    각 청크의 "fragments"는 pipeline.py의 build_chunks와 동일하게 owner당 "청크와 겹친
    부분만"(smart는 owner를 안 쪼개므로 전체 텍스트) - 겹친 owner에 전체 텍스트를 넣으면
    번역이 중복되는 버그(★ 실측 확인, 2026-09-27)를 dry-run 쪽도 똑같이 피해야 한다."""
    smart_groups = group_elements(elements, embeddings._client, count_tokens, min_tokens, max_tokens, window)
    n = len(smart_groups)
    for c in smart_groups:
        c["section"] = section_for_owners(elements, c["owners"])
        c["fragments"] = [(o, get_target(elements, o)["text"]) for o in c["owners"]]

    fixed_spans = make_fixed_chunks(flat_text, n)
    fixed_chunks = [_fragment_chunk(elements, flat_text, s, e) for s, e in fixed_spans]

    sem_spans = make_semantic_chunks(flat_text, embeddings, n)
    sem_missing = sum(1 for sp in sem_spans if sp is None)
    sem_chunks = [_fragment_chunk(elements, flat_text, s, e) for sp in sem_spans if sp for s, e in [sp]]

    return {"smart": (smart_groups, {}), "fixed": (fixed_chunks, {}),
            "semantic": (sem_chunks, {"span_missing": sem_missing})}


def all_translatable_owners(elements):
    owners = set()
    for e in elements:
        if e.get("skip"):
            continue
        if e["label"] == "table":
            owners |= {("cell", e["id"], c["cell_id"]) for c in e["cells"] if not c.get("skip")}
        else:
            owners.add(("el", e["id"]))
    return owners


def strip_check(translations):
    """«» 뗀 결과가 원문과 같은지(중복·누락 0건 확인) - 표시 문자를 전부 지우고 비교."""
    mismatches = 0
    for owner, r in translations.items():
        if MARK_RE.sub("", r["translated"]) != r["orig"]:
            mismatches += 1
    return mismatches


def structure_signature(docx_path):
    from docx import Document as DocxDocument
    doc = DocxDocument(str(docx_path))
    return len(doc.paragraphs), len(doc.tables)


def drawing_signature(pdf_path):
    import pymupdf
    pdf = pymupdf.open(str(pdf_path))
    n_img = sum(len(page.get_images()) for page in pdf)
    n_draw = sum(len(page.get_drawings()) for page in pdf)
    pdf.close()
    return n_img, n_draw


# ---------------------------------------------------------------------------
# docx / pdf 공통 문서 처리
# ---------------------------------------------------------------------------
def dryrun_doc(src_path, fmt, parse_fn, embeddings, count_tokens, call_llm):
    t = {}
    row = {"file": src_path.name}

    t0 = time.monotonic()
    elements = cached_parse(src_path, fmt, parse_fn)
    flat_text = add_spans(elements)
    t["parse"] = time.monotonic() - t0

    t0 = time.monotonic()
    chunks_by_method = build_all_chunks(elements, flat_text, embeddings, count_tokens)
    t["chunk"] = time.monotonic() - t0

    all_owners = all_translatable_owners(elements)
    row["n_smart"] = len(chunks_by_method["smart"][0])
    row["chunks"] = {}
    row["translate"] = {}

    for method, (chunks, extra) in chunks_by_method.items():
        toks = [count_tokens(c["text"]) for c in chunks] or [0]
        covered = set()
        for c in chunks:
            covered.update(o for o in c["owners"] if not is_skip_owner(elements, o))
        row["chunks"][method] = {"n": len(chunks), "tok_min": min(toks), "tok_med": sorted(toks)[len(toks) // 2],
                                  "tok_max": max(toks), "uncovered_owners": len(all_owners - covered), **extra}

        t0 = time.monotonic()
        translations, stats = translate_elements(elements, chunks, call_llm, PROMPT_TEMPLATES[fmt])
        stats["strip_mismatch"] = strip_check(translations)
        stats["time"] = time.monotonic() - t0
        row["translate"][method] = stats

        if method == "smart":
            t0 = time.monotonic()
            out_dir = PIPELINE_RESULTS_DIR / fmt
            out_dir.mkdir(parents=True, exist_ok=True)
            out_path = out_dir / f"{src_path.stem}.dryrun{src_path.suffix}"
            if fmt == "docx":
                before = structure_signature(src_path)
                merge_stats = merge_docx(src_path, elements, translations, out_path)
                after = structure_signature(out_path)
                merge_stats["struct_preserved"] = (before == after)
                merge_stats["unwritten"] = stats["owners_translated"] - merge_stats["written"] - merge_stats["skipped"]
            else:
                before = drawing_signature(src_path)
                merge_stats = merge_pdf(src_path, elements, translations, out_path)
                after = drawing_signature(out_path)
                merge_stats["drawing_preserved"] = (before == after)
                merge_stats["fail_rate"] = merge_stats["failed"] / max(1, merge_stats["written"] + merge_stats["failed"])
            merge_stats["time"] = time.monotonic() - t0
            row["merge"] = merge_stats

    row["time"] = t
    return row


def run_track(files, fmt, parse_fn, embeddings, count_tokens, call_llm, log_path):
    results = []
    for i, f in enumerate(files, 1):
        try:
            row = dryrun_doc(f, fmt, parse_fn, embeddings, count_tokens, call_llm)
            results.append(row)
            print(f"[{fmt} {i}/{len(files)}] {f.name}: "
                  f"N={row['n_smart']} 시간(파싱/청킹)={row['time']['parse']:.1f}/{row['time']['chunk']:.1f}초",
                  flush=True)
        except Exception as e:
            results.append({"file": f.name, "error": str(e)})
            print(f"[{fmt} {i}/{len(files)}] {f.name}: 예외 - {e}", flush=True)
            traceback.print_exc()
        if i % 5 == 0 or i == len(files):
            log_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    log_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    return results


# ---------------------------------------------------------------------------
# srt - mapper.py primitives 재사용(pipeline.py의 process_srt와 같은 방식, 세 청커 비교용으로 확장)
# ---------------------------------------------------------------------------
def dryrun_srt_movie(movie_name, embeddings, bge_tokenizer, srt_call_llm):
    """srt_call_llm: pipeline.setup()이 준 srt 전용 caller - mapper.py의 자기 태그
    규칙(단일 괄호 [n], SRT_MARK_RE, parse_marked)을 그대로 쓴다(pipeline.py 자체의
    [[n]]/parse_numbered와는 다른 형식 - ★ 실측 확인, 2026-09-27: 섞어 쓰면 srt 프롬프트의
    [n]을 못 찾아 missing_numbers가 거의 전부로 나오는 버그가 있었다, pipeline.py의
    fake_llm_srt docstring 참고)."""
    from srt.loader import load_srt
    from srt.mapper import MARK_RE as SRT_MARK_RE
    from srt.mapper import build_prompt, merge_to_units, parse_marked
    from srt.subtitle_pipeline import default_subtitle_chunkers

    en_name, ko_name = SUBTITLE_DATASETS[movie_name]
    src_doc = load_srt(str(SRT_ENG_DIR / en_name))
    chunkers = default_subtitle_chunkers(embeddings._client, bge_tokenizer, "en")

    row = {"movie": movie_name, "chunks": {}, "translate": {}}
    for method, splitter in chunkers.items():
        t0 = time.monotonic()
        chunks = splitter.split_documents([src_doc])
        t_chunk = time.monotonic() - t0
        toks = [len(bge_tokenizer.encode(c.page_content, add_special_tokens=False)) for c in chunks] or [0]
        row["chunks"][method] = {"n": len(chunks), "tok_min": min(toks), "tok_med": sorted(toks)[len(toks) // 2],
                                  "tok_max": max(toks)}

        t0 = time.monotonic()
        n_calls = missing = strip_mismatch = 0
        all_pieces = []
        for chunk in chunks:
            frs, prompt = build_prompt(src_doc, chunk,
                                        instruction="Translate each line to natural Korean. Keep the [n] numbers exactly:\n{body}")
            if not frs:
                continue
            response = srt_call_llm(prompt)
            n_calls += 1
            got_nums = {int(m.group(1)) for m in SRT_MARK_RE.finditer(response)}
            pieces = parse_marked(response, frs)
            for n, ((j, text), (_, s, e, whole)) in enumerate(zip(pieces, frs)):
                if (n + 1) not in got_nums:
                    missing += 1
                elif MARK_RE.sub("", text) != src_doc.page_content[s:e]:
                    strip_mismatch += 1
            all_pieces += pieces
        merge_to_units(src_doc, all_pieces)  # 유닛 병합까지 실제로 도는지 확인(결과 파일 저장은 안 함)
        row["translate"][method] = {"llm_calls": n_calls, "missing_numbers": missing,
                                     "strip_mismatch": strip_mismatch, "time": time.monotonic() - t0}
    return row


def run_srt(embeddings, bge_tokenizer, srt_call_llm, log_path):
    results = []
    movies = list(SUBTITLE_DATASETS)
    for i, movie in enumerate(movies, 1):
        try:
            row = dryrun_srt_movie(movie, embeddings, bge_tokenizer, srt_call_llm)
            results.append(row)
            print(f"[srt {i}/{len(movies)}] {movie}: 완료", flush=True)
        except Exception as e:
            results.append({"movie": movie, "error": str(e)})
            print(f"[srt {i}/{len(movies)}] {movie}: 예외 - {e}", flush=True)
            traceback.print_exc()
    log_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    return results


if __name__ == "__main__":
    from docx_track.parse import parse_docx
    from pdf_track.parse import parse_pdf

    print("모델 준비 중(fake_llm=True)...", flush=True)
    embeddings, count_tokens, call_llm, bge_tokenizer, srt_call_llm = setup(fake_llm_flag=True, need_llm=False)

    print("\n=== docx 45개 dry-run ===", flush=True)
    run_track(docx_files(), "docx", parse_docx, embeddings, count_tokens, call_llm, DRYRUN_DIR / "dryrun_docx.json")

    print("\n=== pdf 50개 dry-run ===", flush=True)
    run_track(pdf_files(), "pdf", parse_pdf, embeddings, count_tokens, call_llm, DRYRUN_DIR / "dryrun_pdf.json")

    print("\n=== srt 전체 dry-run ===", flush=True)
    run_srt(embeddings, bge_tokenizer, srt_call_llm, DRYRUN_DIR / "dryrun_srt.json")

    print("\n=== 전체 dry-run 완료 ===", flush=True)
