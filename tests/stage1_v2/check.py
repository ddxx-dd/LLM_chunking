"""1차 전처리 v2 표본 결과 점검(tests 전용): out/meta_text_v2 의 문서를 원본·기존 meta_text와 비교하고 검토용 파일을 만든다.

  .venv311/bin/python tests/stage1_v2/check.py

PDF  : 원본 글자 층(PyMuPDF) 단어와 비교 - 원본 재현(원본 단어 중 결과에 있는 비율, 기존 meta_text와 v2), 결과의 원본 일치,
       쪽별로 원본 단어의 30% 미만만 담긴 쪽 수, 블록 종류·수식·살린 목차·옮긴 항목·줄 정보 채움 비율
docx : 기존 meta_text와 블록(종류, 글)이 같은지, python-docx로 읽은 원본 글 단어 재현
출력 : out/check.jsonl(문서별), out/check_summary.md(요약),
       out/review/<데이터셋>/<doc_id>.md(결과를 읽기 쉽게: 블록 종류 표시),
       out/review/pages/<doc_id>_p<쪽>.png + .md(PDF 문서마다 가운데 쪽 하나: 원본 쪽 이미지와 그 쪽 블록)
"""
import json
import re
import statistics
import sys
import unicodedata
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / "src"))

from common.config import META_DIR  # noqa: E402

OUT = HERE / "out"
W = re.compile(r"[0-9a-z가-힣]+")
MARK = {"heading": "## ", "table": "[표]\n", "formula": "[수식] ", "list_item": "- ", "footnote": "[각주] ", "code": "[코드] "}


def wc(t):
    return Counter(W.findall(unicodedata.normalize("NFKC", t).lower()))


def ratio(a, b):
    return sum(min(v, b[k]) for k, v in a.items()) / max(1, sum(a.values()))


def docx_text(path):
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def review_text(blocks):
    out = []
    for b in blocks:
        tag = MARK.get(b["type"], "")
        if b.get("role") == "caption":
            tag = "[캡션] "
        out.append(f"{tag}{b['text']}")
    return "\n\n".join(out)


def main():
    import pymupdf

    rows = []
    (OUT / "review" / "pages").mkdir(parents=True, exist_ok=True)
    for f in sorted(OUT.glob("*/meta_text_v2/*.json")):
        ds, doc_id = f.parent.parent.name, f.stem
        new = json.loads(f.read_text(encoding="utf-8"))
        old = json.loads((META_DIR / ds / f"{doc_id}.json").read_text(encoding="utf-8"))
        nb, ob = new["blocks"], old["blocks"]
        (OUT / "review" / ds).mkdir(parents=True, exist_ok=True)
        (OUT / "review" / ds / f"{doc_id}.md").write_text(f"# {ds}/{doc_id}\n\n" + review_text(nb), encoding="utf-8")
        nt = "\n".join(b["text"] for b in nb)
        row = {"dataset": ds, "doc_id": doc_id, "fmt": new["fmt"], "blocks_old": len(ob), "blocks_new": len(nb),
               "types": dict(Counter(b["type"] for b in nb)), "seconds": new["log"].get("seconds")}
        src = ROOT / new["source_file"]
        if new["fmt"] == "pdf":
            with pymupdf.open(str(src)) as d:
                pages = [p.get_text() for p in d]
                raw = wc("\n".join(pages))
                paras = [b for b in nb if b["type"] == "paragraph"]
                low = 0
                for i, p in enumerate(pages, 1):
                    rw = wc(p)
                    on_page = [b["text"] for b in nb if i in ([x[0] for x in b.get("boxes", [])] or [b.get("page")])]   # 쪽을 넘는 블록은 걸친 쪽 모두에
                    if sum(rw.values()) >= 150 and ratio(rw, wc("\n".join(on_page))) < 0.3:
                        low += 1
                body_pages = sorted(p for b in paras for p in [b.get("page")] if p)
                pg = body_pages[len(body_pages) // 2] if body_pages else 1
                d[pg - 1].get_pixmap(dpi=70).save(str(OUT / "review" / "pages" / f"{doc_id}_p{pg}.png"))
            (OUT / "review" / "pages" / f"{doc_id}_p{pg}.md").write_text(
                f"# {ds}/{doc_id} {pg}쪽 (원본 이미지: {doc_id}_p{pg}.png)\n\n" + review_text([b for b in nb if b.get("page") == pg]), encoding="utf-8")
            row.update(pages=len(pages), recall_old=round(ratio(raw, wc("\n".join(b["text"] for b in ob))), 3),
                       recall_new=round(ratio(raw, wc(nt)), 3), precision_new=round(ratio(wc(nt), raw), 3), low_pages=low,
                       formulas=sum(b["type"] == "formula" for b in nb), index_as_body=new["log"].get("index_as_body", 0),
                       trailing_moved=new["log"].get("trailing_moved", 0),
                       geom_filled=round(sum("geom" in b for b in paras) / max(1, len(paras)), 3), review_page=pg)
        else:
            same = [(b["type"], b["text"]) for b in nb] == [(b["type"], b["text"]) for b in ob]
            raw = wc(docx_text(src))
            row.update(same_as_old=same, recall_new=round(ratio(raw, wc(nt)), 3), precision_new=round(ratio(wc(nt), raw), 3))
        rows.append(row)
    with open(OUT / "check.jsonl", "w", encoding="utf-8") as fh:
        for r in rows:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")

    lines = ["# 1차 전처리 v2 표본 점검", ""]
    errors = (OUT / "errors.jsonl").read_text(encoding="utf-8").strip().split("\n") if (OUT / "errors.jsonl").exists() else []
    lines.append(f"변환 오류 {len([e for e in errors if e])}건")
    for ds in ("vectara", "allganize"):
        r = [x for x in rows if x["dataset"] == ds]
        if not r:
            continue
        med = lambda k: statistics.median(x[k] for x in r)
        lines += ["", f"## {ds} ({len(r)}편)",
                  f"- 원본 재현 중앙값: 기존 {med('recall_old'):.1%} → v2 {med('recall_new'):.1%} / 결과의 원본 일치 중앙값 {med('precision_new'):.1%}",
                  f"- 원본 단어 30% 미만만 담긴 쪽: 합계 {sum(x['low_pages'] for x in r)} (쪽 {sum(x['pages'] for x in r)}개 중)",
                  f"- 수식 {sum(x['formulas'] for x in r)}개, 살린 목차 항목 {sum(x['index_as_body'] for x in r)}, 옮긴 항목 {sum(x['trailing_moved'] for x in r)}, "
                  f"줄 정보 채움(문단) 중앙값 {med('geom_filled'):.1%}, 문서당 변환 시간 중앙값 {med('seconds')}초",
                  "", "| 문서 | 쪽 | 재현 기존→v2 | 원본 일치 | 30% 미만 쪽 | 수식 | 살린 목차 | 옮김 | 줄 정보 |", "|---|---|---|---|---|---|---|---|---|"]
        for x in sorted(r, key=lambda x: x["recall_new"] - x["recall_old"], reverse=True):
            lines.append(f"| {x['doc_id'][:34]} | {x['pages']} | {x['recall_old']:.0%}→{x['recall_new']:.0%} | {x['precision_new']:.0%} | "
                         f"{x['low_pages']} | {x['formulas']} | {x['index_as_body']} | {x['trailing_moved']} | {x['geom_filled']:.0%} |")
    r = [x for x in rows if x["fmt"] == "docx"]
    if r:
        lines += ["", f"## docx_corpus ({len(r)}편)",
                  f"- 기존 meta_text와 블록이 완전히 같음: {sum(x['same_as_old'] for x in r)}편",
                  f"- 원본(python-docx) 단어 재현 중앙값 {statistics.median(x['recall_new'] for x in r):.1%}, 최소 {min(x['recall_new'] for x in r):.1%}"]
        for x in r:
            if not x["same_as_old"] or x["recall_new"] < 0.9:
                lines.append(f"  - {x['doc_id'][:50]}: 같음 {x['same_as_old']}, 블록 {x['blocks_old']}→{x['blocks_new']}, 재현 {x['recall_new']:.0%}")
    (OUT / "check_summary.md").write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    main()
